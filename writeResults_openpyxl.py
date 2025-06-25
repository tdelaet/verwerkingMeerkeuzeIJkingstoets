# -*- coding: utf-8 -*-
"""
Created on Wed Jun 18 18:41:40 2025

@author: u0046457
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import numpy as np

# Drempelwaarden
tresholdCorrectAnswer = 0.35
tresholdWrongAnswer = 0.35
tresholdUpperCorrectAnswer = 0.35

tresholdDiffUpperLowerThree = 0.5
tresholdDiffUpperLowerTwo = 0.3
tresholdDiffUpperLowerOne = 0.15

tresholdDifficultyQuestionZero = 0.9
tresholdDifficultyQuestionOne = 0.75
tresholdDifficultyQuestionTwo = 0.5
tresholdDifficultyQuestionThree = 0.25

# Stijlen
font_bold = Font(bold=True)
font_red = Font(color="FF0000")
font_italic = Font(italic=True)

align_horizcenter = Alignment(horizontal="center")
align_horizright = Alignment(horizontal="right")
align_horizleft = Alignment(horizontal="left")
align_vertcenter = Alignment(vertical="center")
align_horizvertcenter = Alignment(horizontal="center", vertical="center")
align_rotated = Alignment(textRotation=90)

border_medium = Side(style="medium")
border_thick = Side(style="thick")
border_bottom_medium = Border(bottom=border_medium)
border_top_medium = Border(top=border_medium)
border_right_medium = Border(right=border_medium)
border_left_medium = Border(left=border_medium)
border_leftbottom_medium = Border(left=border_medium, bottom=border_medium)
border_lefttop_medium = Border(left=border_medium, top=border_medium)
border_righttop_medium = Border(right=border_medium, top=border_medium)
border_rightbottom_medium = Border(right=border_medium, bottom=border_medium)
border_all_medium = Border(left=border_medium, right=border_medium, top=border_medium, bottom=border_medium)
border_all_thick = Border(left=border_thick, right=border_thick, top=border_thick, bottom=border_thick)


pattern_solid_grey = PatternFill("solid", fgColor="C0C0C0")

# Gecombineerde stijlen als dictionaries
style_title = {"font": font_bold, "border": border_all_medium, "alignment": align_horizvertcenter}
style_header = {"font": font_bold, "border": border_bottom_medium, "alignment": align_horizvertcenter}
style_header_borderAll= {"font": font_bold, "border": border_all_medium, "alignment": align_horizvertcenter}
style_header_borderRight = {"font": font_bold, "border": Border(bottom=border_medium, right=border_medium), "alignment": align_horizvertcenter}
style_header_borderRight = {"font": font_bold, "border": Border(bottom=border_medium, right=border_medium), "alignment": align_horizvertcenter}
style_correctAnswer = {"fill": pattern_solid_grey, "font": font_italic}
style_specialAttention = {"font": font_red}
style_border_header = {"border": border_all_thick}
style_border_righttop_medium ={"font": font_bold, "border": border_righttop_medium , "alignment": align_horizvertcenter}
style_bordertop_medium={"font": font_bold, "border": border_lefttop_medium , "alignment": align_horizvertcenter}

# Helper function to apply styles

def apply_style(cell, style):
    for key, value in style.items():
        if isinstance(value, (Font, Alignment, Border, PatternFill)): 
            setattr(cell, key, value) 
        else: 
            setattr(cell, key, value)



# write_results functie (nog afhankelijk van andere functies)
def write_results(outputbook, outputbookperm, numQuestions, correctAnswers, alternatives, blankAnswer,
                  maxTotalScore, content, content_colNrs,
                  columnSeries, deelnemers,
                  numParticipants,
                  totalScore, percentagePass,
                  scoreQuestionsIndicatedSeries,
                  totalScoreDifferentPermutations,
                  medianScore,
                  standardDeviation,
                  averageScore, averageScoreUpper, averageScoreMiddle, averageScoreLower,
                  averageScoreQuestions, averageScoreQuestionsUpper, averageScoreQuestionsMiddle, averageScoreQuestionsLower,
                  averageScoreQuestionsDifferentSeries,
                  numUpper, numMiddle, numLower,
                  numParticipantsSeries,
                  averageScoreSeries, medianScoreSeries, standardDeviationSeries, percentagePassSeries,
                  numQuestionsAlternatives, numQuestionsAlternativesUpper, numQuestionsAlternativesMiddle, numQuestionsAlternativesLower,
                  nameQuestions, classificationQuestionsMod, categoriesQuestions,
                  bordersDistributionStudentsLow, bordersDistributionStudentsHigh, distributionStudentsLow, distributionStudentsHigh):

    write_scoreAllPermutations(outputbookperm, 'ScoreVerschillendeSeries', numParticipants, deelnemers, numQuestions, content, content_colNrs, totalScore, totalScoreDifferentPermutations, columnSeries)
    write_overallStatistics(outputbook, 'GlobaleParameters', totalScore, averageScore, medianScore, standardDeviation, percentagePass, numParticipantsSeries, averageScoreSeries, medianScoreSeries, standardDeviationSeries, percentagePassSeries, maxTotalScore)
    write_averageScoreQuestions(outputbook, 'GemiddeldeScoreVraag', numQuestions, averageScore, averageScoreUpper, averageScoreMiddle, averageScoreLower, averageScoreQuestions, averageScoreQuestionsUpper, averageScoreQuestionsMiddle, averageScoreQuestionsLower, averageScoreSeries, averageScoreQuestionsDifferentSeries, nameQuestions, categoriesQuestions)
    write_percentageAlternativesQuestions(outputbook, "PercentageAlternatieven", numQuestions, correctAnswers, alternatives, blankAnswer, numQuestionsAlternatives, numParticipants, nameQuestions, categoriesQuestions)
    write_percentageAlternativesQuestionsUML(outputbook, "PercentageAlternatievenUML", numQuestions, correctAnswers, alternatives, blankAnswer, numQuestionsAlternativesUpper, numQuestionsAlternativesMiddle, numQuestionsAlternativesLower, numUpper, numMiddle, numLower, nameQuestions, categoriesQuestions)
    questionClassification = write_histogramQuestions(outputbook, "HistogramVragen", numQuestions, scoreQuestionsIndicatedSeries, averageScoreQuestions, nameQuestions, classificationQuestionsMod, categoriesQuestions)
    #write_distributionStudents(outputbookParticipants, bordersDistributionStudentsLow, bordersDistributionStudentsHigh, distributionStudentsLow, distributionStudentsHigh)
    write_distributionStudents(outputbook,"HistogramStudenten",numParticipants,bordersDistributionStudentsLow,bordersDistributionStudentsHigh,distributionStudentsLow,distributionStudentsHigh)


def write_scoreAllPermutations(outputbookperm_loc, nameSheet_loc, numParticipants_loc, deelnemers_loc, numQuestion_loc, content_loc, content_colNrs_loc, totalScore_loc, totalScoreDifferentPermutations_loc, columnSeries_loc):
    sheetC = outputbookperm_loc.create_sheet(title=nameSheet_loc)

    columnCounter = 1
    rowCounter = 1

    # Titel
    sheetC.merge_cells(start_row=rowCounter, start_column=columnCounter, end_row=rowCounter, end_column=columnCounter+8)
    cell = sheetC.cell(row=rowCounter, column=columnCounter, value="Score deelnemers voor verschillende reeksen")
    apply_style(cell, style_title)
    rowCounter += 1

    numSeries_loc = len(totalScoreDifferentPermutations_loc[0])

    # Deelnemersnummers
    sheetC.cell(row=rowCounter, column=columnCounter, value="ijkID")
    apply_style(sheetC.cell(row=rowCounter, column=columnCounter), style_header_borderRight)
    rowCounter += 1
    for i in range(len(deelnemers_loc)):
        cell = sheetC.cell(row=rowCounter, column=columnCounter, value=deelnemers_loc[i])
        apply_style(cell, {"font": font_bold, "border": border_right_medium})
        rowCounter += 1
    columnCounter += 1

    rowCounter = 2
    # Totale score voor aangeduide reeks
    sheetC.cell(row=rowCounter, column=columnCounter, value="aangeduide reeks")
    apply_style(sheetC.cell(row=rowCounter, column=columnCounter), style_header_borderRight)
    rowCounter += 1
    for i in range(len(totalScore_loc)):
        cell = sheetC.cell(row=rowCounter, column=columnCounter, value=totalScore_loc[i])
        apply_style(cell, {"border": border_right_medium})
        rowCounter += 1
    columnCounter += 1

    # Scores voor verschillende reeksen
    for serie in range(1, numSeries_loc + 1):
        rowCounter = 2
        sheetC.cell(row=rowCounter, column=columnCounter, value="reeks " + str(serie))
        apply_style(sheetC.cell(row=rowCounter, column=columnCounter), {"font": font_bold, "border": border_bottom_medium})
        rowCounter += 1
        totalScoreSerie = totalScoreDifferentPermutations_loc[:, serie - 1]
        for i in range(len(totalScore_loc)):
            cell = sheetC.cell(row=rowCounter, column=columnCounter, value=totalScoreSerie[i])
            if serie == columnSeries_loc[i]:
                apply_style(cell, style_correctAnswer)
            elif totalScoreSerie[i] > totalScore_loc[i]:
                apply_style(cell, style_specialAttention)
            rowCounter += 1
        columnCounter += 1

def write_overallStatistics(outputbook_loc, nameSheet_loc, totalScore_loc, averageScore_loc, medianScore_loc,
                            standardDeviation_loc, percentagePass_loc, numParticipantsSeries_loc,
                            averageScoreSeries_loc, medianScoreSeries_loc, standardDeviationSeries_loc,
                            percentagePassSeries_loc, maxTotalScore_loc):
    sheetC = outputbook_loc.create_sheet(title=nameSheet_loc)

    columnCounter = 1
    rowCounter = 1

    # Titel
    sheetC.merge_cells(start_row=rowCounter, start_column=columnCounter, end_row=rowCounter, end_column=columnCounter+8)
    cell = sheetC.cell(row=rowCounter, column=columnCounter, value="Globale statistiek")
    apply_style(cell, style_title)
    rowCounter += 1

    numParticipants_loc = len(totalScore_loc)

    # Aantal deelnemers
    sheetC.cell(row=rowCounter, column=1, value="aantal deelnemers")
    apply_style(sheetC.cell(row=rowCounter, column=1), {"font": font_bold})
    sheetC.cell(row=rowCounter, column=2, value=numParticipants_loc)
    rowCounter += 1

    # Gemiddelde score
    sheetC.cell(row=rowCounter, column=1, value="gemiddelde score")
    apply_style(sheetC.cell(row=rowCounter, column=1), {"font": font_bold})
    sheetC.cell(row=rowCounter, column=2, value=round(averageScore_loc, 2))
    rowCounter += 1

    # Mediaan
    sheetC.cell(row=rowCounter, column=1, value="mediaan")
    apply_style(sheetC.cell(row=rowCounter, column=1), {"font": font_bold})
    sheetC.cell(row=rowCounter, column=2, value=round(medianScore_loc, 2))
    rowCounter += 1

    # Standaard deviatie
    sheetC.cell(row=rowCounter, column=1, value="standaard deviatie")
    apply_style(sheetC.cell(row=rowCounter, column=1), {"font": font_bold})
    sheetC.cell(row=rowCounter, column=2, value=round(standardDeviation_loc, 2))
    rowCounter += 1

    # Percentage geslaagd
    sheetC.cell(row=rowCounter, column=1, value="% geslaagd")
    apply_style(sheetC.cell(row=rowCounter, column=1), {"font": font_bold})
    sheetC.cell(row=rowCounter, column=2, value=round(percentagePass_loc, 2))
    rowCounter += 5

    # Titel voor reeksen
    sheetC.merge_cells(start_row=rowCounter, start_column=1, end_row=rowCounter, end_column=9)
    cell = sheetC.cell(row=rowCounter, column=1, value="Globale statistiek verschillende reeksen")
    apply_style(cell, style_title)
    rowCounter += 1

    # Statistieken per reeks
    for serie in range(len(numParticipantsSeries_loc)):
        sheetC.merge_cells(start_row=rowCounter, start_column=1, end_row=rowCounter, end_column=2)
        cell = sheetC.cell(row=rowCounter, column=1, value=f"serie {serie + 1}")
        apply_style(cell, style_header)
        rowCounter += 1

        labels = ["aantal deelnemers", "gemiddelde score", "mediaan", "standaard deviatie", "% geslaagd"]
        values = [
            numParticipantsSeries_loc[serie],
            round(averageScoreSeries_loc[serie], 2),
            round(medianScoreSeries_loc[serie], 2),
            round(standardDeviationSeries_loc[serie], 2),
            round(percentagePassSeries_loc[serie], 2)
        ]

        for label, value in zip(labels, values):
            sheetC.cell(row=rowCounter, column=1, value=label)
            apply_style(sheetC.cell(row=rowCounter, column=1), {"font": font_bold})
            sheetC.cell(row=rowCounter, column=2, value=value)
            rowCounter += 1

        rowCounter += 1

def write_averageScoreQuestions(outputbook_loc, nameSheet_loc, numQuestions_loc, averageScore_loc, averageScoreUpper_loc,
                                averageScoreMiddle_loc, averageScoreLower_loc, averageScoreQuestions_loc,
                                averageScoreQuestionsUpper_loc, averageScoreQuestionsMiddle_loc, averageScoreQuestionsLower_loc,
                                averageScoreSeries_loc, averageScoreQuestionsDifferentSeries_loc,
                                nameQuestions_loc, categoriesQuestions_loc):
    numSeries = len(averageScoreQuestionsDifferentSeries_loc[0])
    sheetC = outputbook_loc.create_sheet(title=nameSheet_loc)

    columnCounter = 1
    rowCounter = 1

    # Titel
    sheetC.merge_cells(start_row=rowCounter, start_column=columnCounter, end_row=rowCounter, end_column=columnCounter+8)
    cell = sheetC.cell(row=rowCounter, column=columnCounter, value="Gemiddelde score per vraag")
    apply_style(cell, style_title)
    rowCounter += 1

    # Kolommen bovenaan
    columnCounter = 2
    sheetC.cell(row=rowCounter+1, column=columnCounter, value="all")
    apply_style(sheetC.cell(row=rowCounter+1, column=columnCounter), style_header_borderRight)
    columnCounter += 1

    sheetC.merge_cells(start_row=rowCounter, start_column=columnCounter, end_row=rowCounter, end_column=columnCounter+2)
    cell = sheetC.cell(row=rowCounter, column=columnCounter, value="UML")
    apply_style(cell, style_title)
    sheetC.merge_cells(start_row=rowCounter, start_column=columnCounter+3, end_row=rowCounter, end_column=columnCounter+6)
    cell = sheetC.cell(row=rowCounter, column=columnCounter+3, value="reeksen")
    apply_style(cell, style_title)
    rowCounter += 1

    sheetC.cell(row=rowCounter, column=columnCounter, value="upper")
    apply_style(sheetC.cell(row=rowCounter, column=columnCounter), style_header)
    columnCounter += 1
    sheetC.cell(row=rowCounter, column=columnCounter, value="middle")
    apply_style(sheetC.cell(row=rowCounter, column=columnCounter), style_header)
    columnCounter += 1
    sheetC.cell(row=rowCounter, column=columnCounter, value="lower")
    apply_style(sheetC.cell(row=rowCounter, column=columnCounter), style_header_borderRight)
    columnCounter += 1

    for serie in range(numSeries):
        sheetC.cell(row=rowCounter, column=columnCounter, value=f"reeks {serie+1}")
        apply_style(sheetC.cell(row=rowCounter, column=columnCounter), style_header)
        columnCounter += 1

    sheetC.cell(row=rowCounter, column=columnCounter, value="ID vraag")
    apply_style(sheetC.cell(row=rowCounter, column=columnCounter), style_header_borderRight)
    columnCounter += 1

    sheetC.cell(row=rowCounter, column=columnCounter, value="categorie vraag")
    apply_style(sheetC.cell(row=rowCounter, column=columnCounter), style_header_borderRight)
    columnCounter += 1

    rowCounter += 1
    columnCounter = 1

    for question in range(1, numQuestions_loc + 1):
        sheetC.cell(row=rowCounter, column=columnCounter, value=f"vraag{question}")
        apply_style(sheetC.cell(row=rowCounter, column=columnCounter), {"font": font_bold, "border": border_right_medium})
        columnCounter += 1

        if averageScoreQuestions_loc[question-1] < 0:
            cell = sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScoreQuestions_loc[question-1], 3))
            apply_style(cell, {"font": font_red, "border": border_right_medium})
        else:
            cell = sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScoreQuestions_loc[question-1], 3))
            apply_style(cell, {"border": border_right_medium})
        columnCounter += 1

        if averageScoreQuestionsUpper_loc[question-1] <= averageScoreQuestionsLower_loc[question-1] or averageScoreQuestionsUpper_loc[question-1] <= averageScoreQuestionsMiddle_loc[question-1]:
            apply_style(sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScoreQuestionsUpper_loc[question-1], 3)), style_specialAttention)
            columnCounter += 1
            apply_style(sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScoreQuestionsMiddle_loc[question-1], 3)), style_specialAttention)
            columnCounter += 1
            apply_style(sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScoreQuestionsLower_loc[question-1], 3)), {"font": font_red, "border": border_right_medium})
            columnCounter += 1
        else:
            sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScoreQuestionsUpper_loc[question-1], 3))
            columnCounter += 1
            sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScoreQuestionsMiddle_loc[question-1], 3))
            columnCounter += 1
            apply_style(sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScoreQuestionsLower_loc[question-1], 3)), {"border": border_right_medium})
            columnCounter += 1

        for serie in range(numSeries):
            sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScoreQuestionsDifferentSeries_loc[question-1, serie], 3))
            columnCounter += 1
        
        cell = sheetC.cell(row=rowCounter, column=columnCounter+1, value=nameQuestions_loc[question-1])
        apply_style(cell, {"border": border_left_medium})
        columnCounter += 1
        
        cell = sheetC.cell(row=rowCounter, column=columnCounter+1, value=categoriesQuestions_loc[question-1])
        apply_style(cell, {"border": border_left_medium})
        columnCounter += 1
        
        rowCounter += 1
        columnCounter = 1

   
    # Write data to the worksheet
    columnCounter = 1
    cell = sheetC.cell(row=rowCounter, column=columnCounter, value="totaal")
    apply_style(cell, style_border_header)
    columnCounter += 1
    
    cell = sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScore_loc, 3))
    apply_style(cell, style_border_righttop_medium)
    columnCounter += 1
    
    cell = sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScoreUpper_loc, 3))
    apply_style(cell, style_bordertop_medium)
    columnCounter += 1
    
    cell = sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScoreMiddle_loc, 3))
    apply_style(cell, style_bordertop_medium)
    columnCounter += 1
    
    cell = sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScoreLower_loc, 3))
    apply_style(cell, style_border_righttop_medium)
    columnCounter += 1
    
    for serie in range(1, numSeries + 1):
        cell = sheetC.cell(row=rowCounter, column=columnCounter, value=round(averageScoreSeries_loc[serie-1], 3))
        apply_style(cell, style_bordertop_medium)
        columnCounter += 1
    
        

def write_percentageAlternativesQuestions(outputbook_loc, nameSheet_loc, numQuestions_loc, correctAnswers_loc, alternatives_loc, blankAnswer_loc, numQuestionsAlternatives_loc, numParticipants_loc, nameQuestions_loc, categoriesQuestions_loc):
    wb = outputbook_loc
    ws = wb.create_sheet(nameSheet_loc)

    columnCounter = 0
    rowCounter = 0
    ws.merge_cells(start_row=rowCounter+1, start_column=columnCounter+1, end_row=rowCounter+1, end_column=columnCounter+9)
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="Percentage per alternatief")
    apply_style(cell, style_title)
    rowCounter += 1

    # Write alternative names on top
    columnCounter = 1
    for alternative in alternatives_loc + [blankAnswer_loc]:
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=alternative)
        apply_style(cell, style_header)
        columnCounter += 1

    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="ID vraag")
    apply_style(cell, style_header_borderRight)
    columnCounter += 1

    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="categorie vraag")
    apply_style(cell, style_header_borderRight)
    columnCounter += 1
    rowCounter += 1

    for question in range(1, numQuestions_loc + 1):
        columnCounter = 0
        correctAnswer = correctAnswers_loc[question - 1]
        # Loop over alternatives
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="vraag" + str(question))
        apply_style(cell, {"font": font_bold, "border": border_right_medium})
        columnCounter += 1
        alternativeCounter = 0
        for alternative in alternatives_loc + [blankAnswer_loc]:
            percentage = int(round(numQuestionsAlternatives_loc[question - 1, alternativeCounter] / numParticipants_loc * 100, 0))
            if alternative == correctAnswer:
                if (numQuestionsAlternatives_loc[question - 1, alternativeCounter] / numParticipants_loc < tresholdCorrectAnswer):
                    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=percentage)
                    apply_style(cell, {**style_correctAnswer, **style_specialAttention})
                else:
                    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=percentage)
                    apply_style(cell, style_correctAnswer)
            else:
                if (numQuestionsAlternatives_loc[question - 1, alternativeCounter] / numParticipants_loc > tresholdWrongAnswer):
                    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=percentage)
                    apply_style(cell, style_specialAttention)
                else:
                    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=percentage)
            columnCounter += 1
            alternativeCounter += 1
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=nameQuestions_loc[question - 1])
        apply_style(cell, {"border": border_left_medium})
        columnCounter += 1
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=categoriesQuestions_loc[question - 1])
        apply_style(cell, {"border": border_left_medium})
        columnCounter += 1

        rowCounter += 1


def write_numberAlternativesQuestions(outputbook_loc, nameSheet_loc, numQuestions_loc, correctAnswers_loc, alternatives_loc, blankAnswer_loc, numQuestionsAlternatives_loc, numParticipants_loc, nameQuestions_loc, categoriesQuestions_loc):
    wb = outputbook_loc
    ws = wb.create_sheet(nameSheet_loc)

    columnCounter = 0
    rowCounter = 0
    ws.merge_cells(start_row=rowCounter+1, start_column=columnCounter+1, end_row=rowCounter+1, end_column=columnCounter+9)
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="Aantal per alternatief")
    apply_style(cell, style_title)
    rowCounter += 1

    # Write alternative names on top
    columnCounter = 1
    for alternative in alternatives_loc + [blankAnswer_loc]:
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=alternative)
        apply_style(cell, style_header)
        columnCounter += 1

    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="ID vraag")
    apply_style(cell, style_header_borderRight)
    columnCounter += 1

    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="categorie vraag")
    apply_style(cell, style_header_borderRight)
    columnCounter += 1
    rowCounter += 1

    for question in range(1, numQuestions_loc + 1):
        columnCounter = 0
        correctAnswer = correctAnswers_loc[question - 1]
        # Loop over alternatives
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="vraag" + str(question))
        apply_style(cell, {"font": font_bold, "border": border_right_medium})
        columnCounter += 1
        alternativeCounter = 0
        for alternative in alternatives_loc + [blankAnswer_loc]:
            number = round(numQuestionsAlternatives_loc[question - 1, alternativeCounter], 2)
            if alternative == correctAnswer:
                if (numQuestionsAlternatives_loc[question - 1, alternativeCounter] / numParticipants_loc < tresholdCorrectAnswer):
                    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=number)
                    apply_style(cell, {**style_correctAnswer, **style_specialAttention})
                else:
                    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=number)
                    apply_style(cell, style_correctAnswer)
            else:
                if (numQuestionsAlternatives_loc[question - 1, alternativeCounter] / numParticipants_loc > tresholdWrongAnswer):
                    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=number)
                    apply_style(cell, style_specialAttention)
                else:
                    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=number)
            columnCounter += 1
            alternativeCounter += 1
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=nameQuestions_loc[question - 1])
        apply_style(cell, {"border": border_left_medium})
        columnCounter += 1
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=categoriesQuestions_loc[question - 1])
        apply_style(cell, {"border": border_left_medium})
        columnCounter += 1

        rowCounter += 1


def write_percentageAlternativesQuestionsUML(outputbook_loc, nameSheet_loc, numQuestions_loc, correctAnswers_loc, alternatives_loc, blankAnswer_loc, numQuestionsAlternativesUpper_loc, numQuestionsAlternativesMiddle_loc, numQuestionsAlternativesLower_loc, numUpper_loc, numMiddle_loc, numLower_loc, nameQuestions_loc, categoriesQuestions_loc):
    wb = outputbook_loc
    ws = wb.create_sheet(nameSheet_loc)

    columnCounter = 0
    rowCounter = 0
    ws.merge_cells(start_row=rowCounter+1, start_column=columnCounter+1, end_row=rowCounter+1, end_column=columnCounter+9)
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="Percentage per alternatief UML")
    apply_style(cell, style_title)
    rowCounter += 1

    columnCounter = 1
    for alternative in alternatives_loc + [blankAnswer_loc]:
        ws.merge_cells(start_row=rowCounter+1, start_column=columnCounter+1, end_row=rowCounter+1, end_column=columnCounter+3)
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=alternative)
        apply_style(cell, style_header)
        ws.cell(row=rowCounter+2, column=columnCounter+1, value="upper")
        apply_style(ws.cell(row=rowCounter+2, column=columnCounter+1), style_header)
        ws.cell(row=rowCounter+2, column=columnCounter+2, value="middle")
        apply_style(ws.cell(row=rowCounter+2, column=columnCounter+2), style_header)
        ws.cell(row=rowCounter+2, column=columnCounter+3, value="lower")
        apply_style(ws.cell(row=rowCounter+2, column=columnCounter+3), {**style_header_borderRight, **{"border": border_right_medium}})
        columnCounter += 3

    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="onderscheidend vermogen")
    apply_style(cell, style_header)
    columnCounter += 1

    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="ID vraag")
    apply_style(cell, style_header_borderRight)
    columnCounter += 1

    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="categorie vraag")
    apply_style(cell, style_header_borderRight)
    columnCounter += 1

    rowCounter += 2

    for question in range(1, numQuestions_loc + 1):
        columnCounter = 0
        correctAnswer = correctAnswers_loc[question - 1]
        # Loop over alternatives
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="vraag" + str(question))
        apply_style(cell, {"font": font_bold, "border": border_right_medium})
        columnCounter += 1
        alternativeCounter = 0
        for alternative in alternatives_loc + [blankAnswer_loc]:
            upperPerc = numQuestionsAlternativesUpper_loc[question - 1, alternativeCounter] / numUpper_loc if numUpper_loc != 0 else 0
            middlePerc = numQuestionsAlternativesMiddle_loc[question - 1, alternativeCounter] / numMiddle_loc if numMiddle_loc != 0 else 0
            lowerPerc = numQuestionsAlternativesLower_loc[question - 1, alternativeCounter] / numLower_loc if numLower_loc != 0 else 0
            if alternative == correctAnswer:
                if upperPerc < lowerPerc:
                    apply_style(ws.cell(row=rowCounter+1, column=columnCounter+1, value=int(round(upperPerc * 100, 0))), {**style_correctAnswer, **style_specialAttention})
                    apply_style(ws.cell(row=rowCounter+1, column=columnCounter+2, value=int(round(middlePerc * 100, 0))), {**style_correctAnswer, **style_specialAttention})
                    apply_style(ws.cell(row=rowCounter+1, column=columnCounter+3, value=int(round(lowerPerc * 100, 0))), {**style_correctAnswer, **style_specialAttention, **{"border": border_right_medium}})
                else:
                    if upperPerc < tresholdUpperCorrectAnswer:
                        apply_style(ws.cell(row=rowCounter+1, column=columnCounter+1, value=int(round(upperPerc * 100, 0))), {**style_correctAnswer, **style_specialAttention})
                    else:
                        apply_style(ws.cell(row=rowCounter+1, column=columnCounter+1, value=int(round(upperPerc * 100, 0))), style_correctAnswer)
                    apply_style(ws.cell(row=rowCounter+1, column=columnCounter+2, value=int(round(middlePerc * 100, 0))), style_correctAnswer)
                    apply_style(ws.cell(row=rowCounter+1, column=columnCounter+3, value=int(round(lowerPerc * 100, 0))), {**style_correctAnswer, **{"border": border_right_medium}})
                diffUpperLower = upperPerc - lowerPerc
            else:
                if upperPerc > lowerPerc:
                    apply_style(ws.cell(row=rowCounter+1, column=columnCounter+1, value=int(round(upperPerc * 100, 0))), style_specialAttention)
                    apply_style(ws.cell(row=rowCounter+1, column=columnCounter+2, value=int(round(middlePerc * 100, 0))), style_specialAttention)
                    apply_style(ws.cell(row=rowCounter+1, column=columnCounter+3, value=int(round(lowerPerc * 100, 0))), {**style_specialAttention, **{"border": border_right_medium}})
                else:
                    if numQuestionsAlternativesUpper_loc[question - 1, alternativeCounter] > numQuestionsAlternativesUpper_loc[question - 1, alternatives_loc.index(correctAnswer)]:
                        apply_style(ws.cell(row=rowCounter+1, column=columnCounter+1, value=int(round(upperPerc * 100, 0))), style_specialAttention)
                    else:
                        ws.cell(row=rowCounter+1, column=columnCounter+1, value=int(round(upperPerc * 100, 0)))
                    ws.cell(row=rowCounter+1, column=columnCounter+2, value=int(round(middlePerc * 100, 0)))
                    apply_style(ws.cell(row=rowCounter+1, column=columnCounter+3, value=int(round(lowerPerc * 100, 0))), {"border": border_right_medium})
            columnCounter += 3
            alternativeCounter += 1

        if diffUpperLower >= tresholdDiffUpperLowerThree:
            ws.cell(row=rowCounter+1, column=columnCounter+1, value="++++")
        elif diffUpperLower >= tresholdDiffUpperLowerTwo:
            ws.cell(row=rowCounter+1, column=columnCounter+1, value="+++")
        elif diffUpperLower >= tresholdDiffUpperLowerOne:
            ws.cell(row=rowCounter+1, column=columnCounter+1, value="++")
        elif diffUpperLower > 0:
            ws.cell(row=rowCounter+1, column=columnCounter+1, value="+")
        else:
            ws.cell(row=rowCounter+1, column=columnCounter+1, value="0")
        apply_style(ws.cell(row=rowCounter+1, column=columnCounter+1), {"alignment": align_horizvertcenter})
        columnCounter += 1

        apply_style(ws.cell(row=rowCounter+1, column=columnCounter+1, value=nameQuestions_loc[question - 1]), {"border": border_left_medium})
        columnCounter += 1
        apply_style(ws.cell(row=rowCounter+1, column=columnCounter+1, value=categoriesQuestions_loc[question - 1]), {**{"border": border_left_medium}, **{"border": border_right_medium}})
        columnCounter += 1

        rowCounter += 1

def write_histogramQuestions(outputbook_loc, nameSheet_loc, numQuestions_loc, scoreQuestionsIndicatedSeries_loc, averageScoreQuestions_loc, nameQuestions_loc, classificationQuestionsMod_loc, categoriesQuestions_loc):
    wb = outputbook_loc
    ws = wb.create_sheet(nameSheet_loc)
    
    numParticipants = len(scoreQuestionsIndicatedSeries_loc)
    columnCounter = 0
    rowCounter = 0
    ws.merge_cells(start_row=rowCounter+1, start_column=columnCounter+1, end_row=rowCounter+1, end_column=columnCounter+9)
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="Histogram score vragen")
    apply_style(cell, style_title)
    rowCounter += 1
    
    # Column counter
    columnCounter = 0
    # Gemiddelde verdeling scores per vraag
    possibleScores = np.array([-1.0/4.0, 0.0, 1.0, 1.0+1.0/4.0]) # TODO make parameter
    columnCounter = 1
    for possibleScore in possibleScores[:-1]:
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=possibleScore)
        apply_style(cell, style_header)
        columnCounter += 1
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="gemiddelde")
    apply_style(cell, style_header)
    columnCounter += 1
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="type vraag")
    apply_style(cell, style_header)
    columnCounter += 1
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="voorspelde type vraag")
    apply_style(cell, {**style_header, **{"border": border_left_medium}})
    columnCounter += 1
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="ID vraag")
    apply_style(cell, {**style_header, **{"border": border_right_medium}})
    columnCounter += 1
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="categorie vraag")
    apply_style(cell, {**style_header, **{"border": border_right_medium}})
    columnCounter += 1

    rowCounter += 1    
    questionClassification = []
    matrixClassification = np.array([""]*15, dtype=str).reshape(3, 5)
    matrixClassificationCounter = np.zeros(15).reshape(3, 5)
    for question in range(1, numQuestions_loc + 1):
        columnCounter = 0
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="vraag" + str(question))
        apply_style(cell, {"font": font_bold, "border": border_right_medium})
        hist, bins = np.histogram(scoreQuestionsIndicatedSeries_loc[:, question - 1], bins=possibleScores - 1.0/6.0)
        columnCounter += 1    
        for n in hist:        
            if (hist[0]>hist[len(hist)-1] or hist[0]+hist[1]>hist[len(hist)-1]+hist[len(hist)-2]):# More confident in wrong answer than confident in correct answer
                cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=str(n))
                apply_style(cell, style_specialAttention)
            else:
                cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=str(n))
            columnCounter += 1
        if averageScoreQuestions_loc[question - 1] < 0:
            cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=round(averageScoreQuestions_loc[question - 1], 2))
            apply_style(cell, style_specialAttention)
        else:
            cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=round(averageScoreQuestions_loc[question - 1], 2))
        columnCounter += 1
        correctPerc = float(hist[-1]) / numParticipants
        if correctPerc > tresholdDifficultyQuestionZero:
            cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="0")
            apply_style(cell, {**style_specialAttention, **{"alignment": align_horizvertcenter, "border": border_left_medium}})
            questionClassification.append("0")
            colClass = 0
        elif correctPerc > tresholdDifficultyQuestionOne:
            cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="*")
            apply_style(cell, {"alignment": align_horizvertcenter, "border": border_left_medium})
            questionClassification.append("*")
            colClass = 1
        elif correctPerc > tresholdDifficultyQuestionTwo: 
            cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="**")
            apply_style(cell, {"alignment": align_horizvertcenter, "border": border_left_medium})
            questionClassification.append("**")
            colClass = 2
        elif correctPerc > tresholdDifficultyQuestionThree: 
            cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="***")
            apply_style(cell, {"alignment": align_horizvertcenter, "border": border_left_medium})
            questionClassification.append("***")
            colClass = 3
        else: 
            cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="****")
            apply_style(cell, {**style_specialAttention, **{"alignment": align_horizvertcenter, "border": border_left_medium}})
            questionClassification.append("****")
            colClass = 4
        columnCounter += 1
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=classificationQuestionsMod_loc[question - 1])
        apply_style(cell, {"border": border_right_medium})
        columnCounter += 1        
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=nameQuestions_loc[question - 1])
        apply_style(cell, {"border": border_right_medium})
        columnCounter += 1
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=categoriesQuestions_loc[question - 1])
        apply_style(cell, {"border": border_right_medium})
        columnCounter += 1        

        rowClass = len(classificationQuestionsMod_loc[question - 1]) - 1       
        matrixClassification[rowClass][colClass] = matrixClassification[rowClass][colClass] + str(question) + " "
        matrixClassificationCounter[rowClass][colClass] += 1
        
        rowCounter += 1




def write_scoreStudents(outputbook_loc, nameSheet_loc, permutations_loc, numParticipants_loc, deelnemers_loc, numQuestions_loc, numAlternatives_loc, content_loc, content_colNrs_loc, totalScore_loc, scoreQuestionsIndicatedSeries_loc, columnSeries_loc, matrixAnswers, numberCorrectAnswers_loc, numberWrongAnswers_loc, numberBlankAnswers_loc, numberNeutralizedAnswers_loc):
    wb = outputbook_loc
    ws = wb.create_sheet(nameSheet_loc)

    columnCounter = 0
    rowCounter = 0
    
    # Deelnemersnummers
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="ijkID")
    apply_style(cell, {**style_header, **{"border": border_right_medium}})
    rowCounter += 1
    for i in range(len(deelnemers_loc)):
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=deelnemers_loc[i])
        apply_style(cell, {**style_header, **{"border": border_right_medium}})
        rowCounter += 1
    columnCounter += 1
    
    rowCounter = 0
    # Total score for indicated series
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="totale score")
    apply_style(cell, style_header)
    rowCounter += 1
    for i in range(len(totalScore_loc)):
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=totalScore_loc[i])
        rowCounter += 1
    columnCounter += 1
    
    rowCounter = 0
    # Indicated series
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="reeks")
    apply_style(cell, style_header)
    rowCounter += 1
    for i in range(len(totalScore_loc)):
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=int(columnSeries_loc[i]))
        rowCounter += 1
    columnCounter += 1
    
    rowCounter = 0
    # Number of correct answers (including neutralized answers)
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="aantal juist")
    apply_style(cell, style_header)
    rowCounter += 1
    for i in range(len(numberCorrectAnswers_loc)):
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=numberCorrectAnswers_loc[i] + numberNeutralizedAnswers_loc[i])
        rowCounter += 1
    columnCounter += 1
    
    rowCounter = 0
    # Number of wrong answers
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="aantal fout")
    apply_style(cell, style_header)
    rowCounter += 1
    for i in range(len(numberWrongAnswers_loc)):
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=numberWrongAnswers_loc[i])
        rowCounter += 1
    columnCounter += 1
    
    rowCounter = 0
    # Number of blank answers
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="aantal blanco")
    apply_style(cell, style_header)
    rowCounter += 1
    for i in range(len(numberBlankAnswers_loc)):
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=numberBlankAnswers_loc[i])
        rowCounter += 1
    columnCounter += 1
    
    # Score for different questions
    rowCounter = 0
    columnCounterHeader = columnCounter
    for question in range(1, numQuestions_loc + 1):
        cell = ws.cell(row=rowCounter+1, column=columnCounterHeader+1, value="score vraag " + str(question))
        apply_style(cell, style_header)
        columnCounterHeader += 1
        
    columnCounterScoreQuestions = columnCounter   
    rowCounter = 1
    for participant in range(len(totalScore_loc)): # Loop over participants
        columnCounter = columnCounterScoreQuestions
        score = scoreQuestionsIndicatedSeries_loc[participant, :]
        
        sorted_score = [score[int(i-1)] for i in permutations_loc[int(columnSeries_loc[participant]-1)]]
        for question in range(1, numQuestions_loc + 1):
            ws.cell(row=rowCounter+1, column=columnCounter+1, value=sorted_score[question-1])
            columnCounter += 1
        rowCounter += 1            
 
    # Answer for different questions and alternatives
    for question in range(1, numQuestions_loc + 1):
        rowCounter = 0
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="antwoord vraag " + str(question))
        apply_style(cell, style_header)
        answer = matrixAnswers[:, question-1]
        rowCounter += 1
        for i in range(len(totalScore_loc)):
            ws.cell(row=rowCounter+1, column=columnCounter+1, value=answer[i])
            rowCounter += 1                    
        columnCounter += 1



def write_resultsFile(outputbook_loc, nameSheet_loc, permutations_loc, numParticipants_loc, deelnemers_loc, numQuestions_loc, numAlternatives_loc, content_loc, content_colNrs_loc, totalScore_loc, scoreQuestionsIndicatedSeries_loc, columnSeries_loc, matrixAnswers, numberCorrectAnswers_loc, numberWrongAnswers_loc, numberBlankAnswers_loc, numberNeutralizedAnswers_loc):
    wb = outputbook_loc
    ws = wb.create_sheet(nameSheet_loc)

    columnCounter = 0
    rowCounter = 0
    
    # Deelnemersnummers
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="ijkID")
    apply_style(cell, {**style_header, **{"border": border_right_medium}})
    rowCounter += 1
    for i in range(len(deelnemers_loc)):
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=deelnemers_loc[i])
        apply_style(cell, {**style_header, **{"border": border_right_medium}})
        rowCounter += 1
    columnCounter += 1
    
    rowCounter = 0
    # Total score for indicated series
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="totale score")
    apply_style(cell, style_header)
    rowCounter += 1
    for i in range(len(totalScore_loc)):
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=totalScore_loc[i])
        rowCounter += 1
    columnCounter += 1
    
    rowCounter = 0
    # Indicated series
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="reeks")
    apply_style(cell, style_header)
    rowCounter += 1
    for i in range(len(totalScore_loc)):
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=int(columnSeries_loc[i]))
        rowCounter += 1
    columnCounter += 1
    
    rowCounter = 0
    # Number of correct answers (including neutralized answers)
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="aantal juist")
    apply_style(cell, style_header)
    rowCounter += 1
    for i in range(len(numberCorrectAnswers_loc)):
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=numberCorrectAnswers_loc[i] + numberNeutralizedAnswers_loc[i])
        rowCounter += 1
    columnCounter += 1
    
    rowCounter = 0
    # Number of wrong answers
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="aantal fout")
    apply_style(cell, style_header)
    rowCounter += 1
    for i in range(len(numberWrongAnswers_loc)):
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=numberWrongAnswers_loc[i])
        rowCounter += 1
    columnCounter += 1
    
    rowCounter = 0
    # Number of blank answers
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="aantal blanco")
    apply_style(cell, style_header)
    rowCounter += 1
    for i in range(len(numberBlankAnswers_loc)):
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=numberBlankAnswers_loc[i])
        rowCounter += 1
    columnCounter += 1


def write_scoreCategoriesStudents(outputbook_loc, nameSheet_loc, deelnemers_loc, totalScore_loc, categoriesQuestions_loc, scoreCategories_loc):
    wb = outputbook_loc
    ws = wb.create_sheet(nameSheet_loc)

    columnCounter = 0
    rowCounter = 0
    
    # Deelnemersnummers
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="ijkID")
    apply_style(cell, {**style_header, **{"border": border_right_medium}})
    columnCounter += 1
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="totale score")
    apply_style(cell, {**style_header, **{"border": border_right_medium}})
    rowCounter += 1
    columnCounter = 0
    
    for i in range(len(deelnemers_loc)):
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=deelnemers_loc[i])
        apply_style(cell, {**style_header, **{"border": border_right_medium}})
        cell = ws.cell(row=rowCounter+1, column=columnCounter+2, value=totalScore_loc[i])
        apply_style(cell, {"border": border_right_medium})
        rowCounter += 1
    columnCounter += 2
    
    columnCounterCat = columnCounter
    rowCounter = 0
    
    for categorie in set(categoriesQuestions_loc):
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=categorie)
        apply_style(cell, style_header)
        columnCounter += 1
    rowCounter += 1

    columnCounter = columnCounterCat
    for deelnemer in range(scoreCategories_loc.shape[1]):
        for categorie in range(len(set(categoriesQuestions_loc))):
            ws.cell(row=rowCounter+1, column=columnCounter+1, value=scoreCategories_loc[categorie][deelnemer])
            columnCounter += 1
        rowCounter += 1
        columnCounter = columnCounterCat


def write_overallStatisticsInstellingen(outputbook_loc, nameSheet_loc, instellingen_loc, numParticipants_tot_loc, numParticipants_stacked_tot_loc, averageScore_tot_loc, averageScore_stacked_tot_loc, medianScore_tot_loc, medianScore_stacked_tot_loc, standardDeviation_tot_loc, standardDeviation_stacked_tot_loc, percentagePass_tot_loc, percentagePass_stacked_tot_loc):
    wb = outputbook_loc
    ws = wb.create_sheet(nameSheet_loc)

    columnCounter = 0
    rowCounter = 0
    ws.merge_cells(start_row=rowCounter+1, start_column=columnCounter+1, end_row=rowCounter+1, end_column=columnCounter+9)
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="Globale statistiek")
    apply_style(cell, style_title)
    rowCounter += 1

    columnCounter = 0
    rowCounter = 1
    
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="aantal deelnemers")
    apply_style(cell, {"font": font_bold})
    columnCounter += 1  
    ws.cell(row=rowCounter+1, column=columnCounter+1, value=str(numParticipants_tot_loc))
    rowCounter += 1
    
    columnCounter = 0
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="gemiddelde score")
    apply_style(cell, {"font": font_bold})
    columnCounter += 1  
    ws.cell(row=rowCounter+1, column=columnCounter+1, value=round(averageScore_tot_loc, 2))
    rowCounter += 1
    
    columnCounter = 0
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="mediaan")
    apply_style(cell, {"font": font_bold})
    columnCounter += 1  
    ws.cell(row=rowCounter+1, column=columnCounter+1, value=round(medianScore_tot_loc, 2))
    rowCounter += 1

    columnCounter = 0
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="standaard deviatie")
    apply_style(cell, {"font": font_bold})
    columnCounter += 1  
    ws.cell(row=rowCounter+1, column=columnCounter+1, value=round(standardDeviation_tot_loc, 2))
    rowCounter += 1
        
    columnCounter = 0
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="% geslaagd")
    apply_style(cell, {"font": font_bold})
    columnCounter += 1 
    ws.cell(row=rowCounter+1, column=columnCounter+1, value=round(percentagePass_tot_loc, 2))
    
    rowCounter += 5
    columnCounter = 0
    ws.merge_cells(start_row=rowCounter+1, start_column=columnCounter+1, end_row=rowCounter+1, end_column=columnCounter+9)
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="Globale statistiek verschillende reeksen")
    apply_style(cell, style_title)
    rowCounter += 1
    
    numInstellingen = len(numParticipants_stacked_tot_loc)
    
    for instelling in range(numInstellingen):
        columnCounter = 0        
        ws.merge_cells(start_row=rowCounter+1, start_column=columnCounter+1, end_row=rowCounter+1, end_column=columnCounter+2)
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="instelling " + instellingen_loc[instelling])
        apply_style(cell, style_header)
        rowCounter += 1
        
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="aantal deelnemers")
        apply_style(cell, {"font": font_bold})
        columnCounter += 1  
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=str(numParticipants_stacked_tot_loc[instelling][0]))
        rowCounter += 1
        
        columnCounter = 0
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="gemiddelde score")
        apply_style(cell, {"font": font_bold})
        columnCounter += 1  
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=round(averageScore_stacked_tot_loc[instelling][0], 2))
        rowCounter += 1
    
        
        columnCounter = 0
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="mediaan")
        apply_style(cell, {"font": font_bold})
        columnCounter += 1  
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=round(medianScore_stacked_tot_loc[instelling][0], 2))
        rowCounter += 1
        
        columnCounter = 0
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="standaard deviatie")
        apply_style(cell, {"font": font_bold})
        columnCounter += 1  
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=round(standardDeviation_stacked_tot_loc[instelling][0], 2))
        rowCounter += 1
                
        columnCounter = 0
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="% geslaagd")
        apply_style(cell, {"font": font_bold})
        columnCounter += 1 
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=round(percentagePass_stacked_tot_loc[instelling][0], 2))
        
        rowCounter += 2


def write_distributionStudents(outputbook_loc, nameSheet_loc, numParticpants_loc, bordersDistributionStudentsLow_loc, bordersDistributionStudentsHigh_loc, distributionStudentsLow_loc, distributionStudentsHigh_loc):
    wb = outputbook_loc
    ws = wb.create_sheet(nameSheet_loc)
    columnCounter = 0
    rowCounter = 0
    ws.merge_cells(start_row=rowCounter+1, start_column=columnCounter+1, end_row=rowCounter+1, end_column=columnCounter+9)
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="Histogram studenten")
    apply_style(cell, style_title)
    rowCounter += 2
    columnCounter = 0
        
    ws.merge_cells(start_row=rowCounter+1, start_column=columnCounter+1, end_row=rowCounter+1, end_column=columnCounter+5)
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="percentage met score <= ")
    apply_style(cell, style_title)
    rowCounter += 1
    columnCounter = 0
    
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="score")
    apply_style(cell, {**style_header, **{"border": border_right_medium}})
    columnCounter += 1
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="percentage")
    apply_style(cell, {**style_header, **{"border": border_right_medium}})
    columnCounter += 1
    
    rowCounter += 1
    columnCounter = 0
    counter = 0
    for score in bordersDistributionStudentsLow_loc:
        columnCounter = 0
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=score)
        apply_style(cell, {**style_header, **{"border": border_right_medium}})
        columnCounter += 1
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=round(float(distributionStudentsLow_loc[counter]) / float(numParticpants_loc) * 100, 1))
        columnCounter += 1  
        
        rowCounter += 1
        counter += 1
        
    rowCounter += 2
    columnCounter = 0
    ws.merge_cells(start_row=rowCounter+1, start_column=columnCounter+1, end_row=rowCounter+1, end_column=columnCounter+5)
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="percentage met score >= ")
    apply_style(cell, style_title)
    rowCounter += 1
    columnCounter = 0

    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="score")
    apply_style(cell, {**style_header, **{"border": border_right_medium}})
    columnCounter += 1
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="percentage")
    apply_style(cell, {**style_header, **{"border": border_right_medium}})
    columnCounter += 1
    
    rowCounter += 1
    columnCounter = 0
    counter = 0
    for score in bordersDistributionStudentsHigh_loc:
        columnCounter = 0
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=score)
        apply_style(cell, {**style_header, **{"border": border_right_medium}})
        columnCounter += 1
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=round(float(distributionStudentsHigh_loc[counter]) / float(numParticpants_loc) * 100, 1))
        columnCounter += 1  
        
        rowCounter += 1
        counter += 1

#STILL OLD CODE
def write_feedbackStudents(outputbook_loc,permutations_loc,numParticipants_loc,deelnemers_loc, numQuestions_loc,alternatives_loc,numAlternatives_loc,content_loc,content_colNrs_loc,totalScore_loc,scoreQuestionsIndicatedSeries_loc,columnSeries_loc,matrixAnswers,categorieQuestions_loc,scoreCategories_loc,
                            averageScoreQuestions_tot_loc,averageScoreQuestionsUpper_tot_loc,averageScoreQuestionsMiddle_tot_loc,averageScoreQuestionsLower_tot_loc
                            ,correctAnswers_loc, numQuestionsAlternatives_loc):
    
    print("WARNING write_feedbackStudents not implemented currently in writeResults")
    
#     orderedParticipants = sorted(range(len(deelnemers_loc)), key=lambda k: deelnemers_loc[k])
    
#     for participant in orderedParticipants: # range(len(totalScore_loc)):
#         sheetC = outputbook_loc.add_sheet(str(int(deelnemers_loc[participant])))

#         columnCounter = 0;
#         rowCounter = 0;
        
#         #deelnemersnummers
#         sheetC.write(rowCounter, 0,"ijkID: ", style=easyxf(font_bold + border_lefttop_medium) ) 
#         columnCounter +=1
#         sheetC.write(rowCounter,columnCounter,deelnemers_loc[participant], style=easyxf(font_bold + border_righttop_medium))
        
#         rowCounter+=1
#         columnCounter = 0
        
#         #total score for indicated series
#         sheetC.write(rowCounter,columnCounter,"score",style=easyxf(border_left_medium + font_bold))
#         columnCounter+=1
#         sheetC.write(rowCounter,columnCounter,totalScore_loc[participant],style=easyxf(border_right_medium))
        
#         rowCounter+=1;
#         columnCounter = 0;
        
#         #indicated series
#         sheetC.write(rowCounter,columnCounter,"reeks",style=easyxf(font_bold + border_leftbottom_medium)) 
#         columnCounter+=1
#         sheetC.write(rowCounter,columnCounter,columnSeries_loc[participant],style=easyxf(border_rightbottom_medium))
#         rowCounter+=1;
        
#         rowCounter +=2; 
#         rowBegin2 = rowCounter
        
#         columnCounter=0
#         counterCategorie = 0
#         sheetC.write(rowCounter,columnCounter,"",style=easyxf(border_bottom_medium))
#         columnCounter+=1
#         sheetC.write(rowCounter,columnCounter,"",style=easyxf(border_bottom_medium))
#         rowCounter+=1
#         columnCounter=0
#         for categorie in set(categorieQuestions_loc):
#             sheetC.write(rowCounter,columnCounter,categorie,style=easyxf(font_bold + border_left_medium))
#             columnCounter+=1
#             sheetC.write(rowCounter,columnCounter,scoreCategories_loc[counterCategorie][participant],style=easyxf(font_bold + border_right_medium))
#             rowCounter+=1
#             columnCounter-=1
#             counterCategorie+=1
#         sheetC.write(rowCounter,columnCounter,"",style=easyxf(border_top_medium))
#         columnCounter+=1
#         sheetC.write(rowCounter,columnCounter,"",style=easyxf(border_top_medium))
#         rowCounter+=1
        
#         score = scoreQuestionsIndicatedSeries_loc[participant,:]
#         sorted_score = [score[int(i-1)] for i in permutations_loc[int(columnSeries_loc[participant]-1)]]
        
#         numBlank = sum(score == 0)
#         numCorrect = sum(score == 1.0)
#         numWrong = sum(score == -1.0/(numAlternatives_loc-1))    
        
#         columnOffset = 4
#         columnCounter = columnOffset
#         rowCounter = rowBegin2
#         sheetC.write(rowCounter,columnCounter,"juist",style=easyxf(font_bold + border_lefttop_medium))
#         columnCounter +=1
#         sheetC.write(rowCounter,columnCounter,str(numCorrect),style=easyxf( border_righttop_medium))
#         rowCounter+=1
        
#         columnCounter = columnOffset
#         sheetC.write(rowCounter,columnCounter,"fout",style=easyxf(font_bold+border_left_medium))
#         columnCounter +=1
#         sheetC.write(rowCounter,columnCounter,str(numWrong),style=easyxf(border_right_medium))
#         rowCounter+=1
        
#         columnCounter = columnOffset
#         sheetC.write(rowCounter,columnCounter,"blanco",style=easyxf(font_bold+border_leftbottom_medium))
#         columnCounter +=1
#         sheetC.write(rowCounter,columnCounter,str(numBlank),style=easyxf(border_rightbottom_medium))
#         rowCounter+=1
        
#         rowOffset = rowCounter+len(set(categorieQuestions_loc)); 
        
#         #score for different questions
#         #beware scores are stored per question without the permutation; 
#         #so for the student the scores have to be back-permutated to the order they got
#         rowCounter = rowOffset;
#         columnCounter = 0
        
#         #write heading  
#         sheetC.write(rowCounter,columnCounter,"vraag",style=easyxf(style_header))
#         columnCounter+=1
#         sheetC.write(rowCounter,columnCounter,"score",style=easyxf(style_header))
#         columnCounter+=1        
#         sheetC.write(rowCounter,columnCounter,"antwoord",style=easyxf(style_header))       
#         columnCounter+=1       
#         sheetC.write(rowCounter,columnCounter,"sleutel",style=easyxf(style_header))       
#         columnCounter+=1   
#         sheetC.write(rowCounter,columnCounter,"vraagnr.",style=easyxf(style_header)) 
#         columnCounter+=1        
#         sheetC.write(rowCounter,columnCounter,"type",style=easyxf(style_header)) 
#         columnCounter+=1        
#         sheetC.write(rowCounter,columnCounter,"gem.",style=easyxf(style_header))    
#         columnCounter+=1        
#         sheetC.write(rowCounter,columnCounter,"upper",style=easyxf(style_header))
#         columnCounter+=1        
#         sheetC.write(rowCounter,columnCounter,"lower",style=easyxf(style_header)) 
#         columnCounter+=1        
#         sheetC.write(rowCounter,columnCounter,"% juist",style=easyxf(style_header))
#         columnCounter+=1        
#         sheetC.write(rowCounter,columnCounter,"% blanco",style=easyxf(style_header))

#         rowCounter+=1
        

#         for question in range(1,numQuestions_loc+1):
#             columnCounter = 0
#             questionNumberSerie1 = permutations_loc[int(columnSeries_loc[participant]-1),int(question-1)]
#             correctAnswer = correctAnswers_loc[int(questionNumberSerie1-1)]
#             #print("test")
#             #print(questionNumberSerie1)
#             sheetC.write(rowCounter,columnCounter,str(question),style=easyxf(font_bold + border_right_medium + align_horizright))
#             columnCounter+=1
#             sheetC.write(rowCounter,columnCounter,sorted_score[int(question-1)],style=easyxf(align_horizleft))
#             columnCounter+=1
#             answer = matrixAnswers[participant,question-1]
#             sheetC.write(rowCounter,columnCounter,answer,style=easyxf(align_horizleft))   
#             columnCounter+=1   
#             sheetC.write(rowCounter,columnCounter,correctAnswer,style=easyxf(align_horizleft))   
#             columnCounter+=1  
#             sheetC.write(rowCounter,columnCounter,questionNumberSerie1,style=easyxf(align_horizleft))   
#             columnCounter+=1   
#             sheetC.write(rowCounter,columnCounter,categorieQuestions_loc[int(questionNumberSerie1-1)],style=easyxf(align_horizleft))  
#             columnCounter+=1;    
#             sheetC.write(rowCounter,columnCounter,round(averageScoreQuestions_tot_loc[int(questionNumberSerie1-1)],2),style=easyxf(align_horizleft))  
#             columnCounter+=1;    
#             sheetC.write(rowCounter,columnCounter,round(averageScoreQuestionsUpper_tot_loc[int(questionNumberSerie1-1)],2),style=easyxf(align_horizleft) )
#             columnCounter+=1;    
#             sheetC.write(rowCounter,columnCounter,round(averageScoreQuestionsLower_tot_loc[int(questionNumberSerie1-1)],2),style=easyxf(align_horizleft))             
 
#             percCorrect = int(round(numQuestionsAlternatives_loc[int(questionNumberSerie1-1),alternatives_loc.index(correctAnswer)]/numParticipants_loc*100,0))
#             columnCounter+=1;    
#             sheetC.write(rowCounter,columnCounter,percCorrect,style=easyxf(align_horizleft))
#             percBlank = int(round(numQuestionsAlternatives_loc[int(questionNumberSerie1-1),numAlternatives_loc]/numParticipants_loc*100,0))
#             columnCounter+=1;    
#             sheetC.write(rowCounter,columnCounter,percBlank,style=easyxf(align_horizleft))            
#             rowCounter+=1; 

def write_feedbackPlatform(outputFolder_loc, permutations_loc, numParticipants_loc, deelnemers_loc, numQuestions_loc, alternatives_loc, numAlternatives_loc, content_loc, content_colNrs_loc, totalScore_loc, scoreQuestionsIndicatedSeries_loc, columnSeries_loc, matrixAnswers, categorieQuestions_loc, scoreCategories_loc, averageScoreQuestions_tot_loc, averageScoreQuestionsUpper_tot_loc, averageScoreQuestionsMiddle_tot_loc, averageScoreQuestionsLower_tot_loc, correctAnswers_loc, numQuestionsAlternatives_loc, blankAnswer_loc):
    wb = Workbook()
    ws = wb.active
    ws.title = "Feedback Platform"

    orderedParticipants = sorted(range(len(deelnemers_loc)), key=lambda k: deelnemers_loc[k])

    # Write header
    headers = ["ijkID", "reeks", "score"]
    headers += [f"cat: {categorie}" for categorie in set(categorieQuestions_loc)]
    headers += ["aantal juist", "aantal fout", "aantal blanco"]
    for question in range(1, numQuestions_loc + 1):
        headers += [
            f"vraag{question}: score", f"vraag{question}: antwoord", f"vraag{question}: sleutel",
            f"vraag{question}: type", f"vraag{question}: gem.", f"vraag{question}: upper",
            f"vraag{question}: lower", f"vraag{question}: %juist", f"vraag{question}: %blanco"
        ]
        headers += [f"aantal {alternative}" for alternative in alternatives_loc]
        headers.append(f"aantal {blankAnswer_loc}")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        apply_style(cell, style_header)

    # Write data
    for row_num, participant in enumerate(orderedParticipants, start=2):
        data = [
            int(deelnemers_loc[participant]),
            int(columnSeries_loc[participant]),
            int(totalScore_loc[participant])
        ]
        data += [scoreCategories_loc[counterCategorie][participant] for counterCategorie in range(len(set(categorieQuestions_loc)))]
        
        score = scoreQuestionsIndicatedSeries_loc[participant, :]
        sorted_score = [score[int(i - 1)] for i in permutations_loc[int(columnSeries_loc[participant] - 1)]]
        
        numBlank = sum(score == 0)
        numCorrect = sum(score == 1.0)
        numWrong = sum(score == -1.0 / (numAlternatives_loc - 1))
        
        data += [numCorrect, numWrong, numBlank]

        for question in range(1, numQuestions_loc + 1):
            questionNumberSerie1 = permutations_loc[int(columnSeries_loc[participant] - 1), int(question - 1)]
            correctAnswer = correctAnswers_loc[int(questionNumberSerie1 - 1)]
            answer = matrixAnswers[participant, int(question - 1)]
            percCorrect = int(round(numQuestionsAlternatives_loc[int(questionNumberSerie1 - 1), alternatives_loc.index(correctAnswer)] / numParticipants_loc * 100, 0))
            percBlank = int(round(numQuestionsAlternatives_loc[int(questionNumberSerie1 - 1), numAlternatives_loc] / numParticipants_loc * 100, 0))
            data += [
                sorted_score[int(question - 1)], answer, correctAnswer, categorieQuestions_loc[int(questionNumberSerie1 - 1)],
                round(averageScoreQuestions_tot_loc[int(questionNumberSerie1 - 1)], 2),
                int(round(averageScoreQuestionsUpper_tot_loc[int(questionNumberSerie1 - 1)] * 100, 0)),
                int(round(averageScoreQuestionsLower_tot_loc[int(questionNumberSerie1 - 1)] * 100, 0)),
                percCorrect, percBlank
            ]
            data += [int(round(numQuestionsAlternatives_loc[int(questionNumberSerie1 - 1), alternative], 0)) for alternative in range(numAlternatives_loc)]
            data.append(int(round(numQuestionsAlternatives_loc[int(questionNumberSerie1 - 1), numAlternatives_loc], 0)))

        for col_num, value in enumerate(data, 1):
            ws.cell(row=row_num, column=col_num, value=value)

    wb.save(outputFolder_loc + 'feedbackPlatform.xlsx')
    
    
def write_scoreStudentsNonPermutated(outputbook_loc, nameSheet_loc, permutations_loc, numParticipants_loc, deelnemers_loc, numQuestions_loc, numAlternatives_loc, alternatives_loc, content_loc, content_colNrs_loc, totalScore_loc, scoreQuestionsIndicatedSeries_loc, columnSeries_loc, matrixAnswers):
    wb = outputbook_loc
    ws = wb.create_sheet(title=nameSheet_loc)

    columnCounter = 0
    rowCounter = 0
    
    # Deelnemersnummers
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="studentennummer")
    apply_style(cell, style_header_borderRight)
    rowCounter += 1
    for i in range(len(deelnemers_loc)):
        cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value=deelnemers_loc[i])
        apply_style(cell, style_header_borderRight)
        rowCounter += 1
    columnCounter += 1
    
    rowCounter = 0
    # Total score for indicated series
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="totale score")
    apply_style(cell, style_header)
    rowCounter += 1
    for i in range(len(totalScore_loc)):
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=totalScore_loc[i])
        rowCounter += 1
    columnCounter += 1
    
    rowCounter = 0
    # Indicated series
    cell = ws.cell(row=rowCounter+1, column=columnCounter+1, value="reeks")
    apply_style(cell, style_header)
    rowCounter += 1
    for i in range(len(totalScore_loc)):
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=columnSeries_loc[i])
        rowCounter += 1
    columnCounter += 1
    
    # Score for different questions
    rowCounter = 0
    # Write heading    
    columnCounterHeader = columnCounter
    for question in range(1, numQuestions_loc + 1):
        cell = ws.cell(row=rowCounter+1, column=columnCounterHeader+1, value="score vraag " + str(question))
        apply_style(cell, style_header)
        columnCounterHeader += 1
        
    columnCounterScoreQuestions = columnCounter   
    rowCounter = 1
    for participant in range(len(totalScore_loc)): # Loop over participants
        columnCounter = columnCounterScoreQuestions
        score = scoreQuestionsIndicatedSeries_loc[participant, :]
        
        for question in range(1, numQuestions_loc + 1):
            ws.cell(row=rowCounter+1, column=columnCounter+1, value=score[int(question-1)])
            columnCounter += 1
        rowCounter += 1            
 
    # Answers for different questions
    # Beware answers are stored per question with the permutation; 
    # so for the us the answers have to be permutated 
    # Write heading    
    for question in range(1, numQuestions_loc + 1):
        for alternative in alternatives_loc:
            cell = ws.cell(row=1, column=columnCounterHeader+1, value="antwoord vraag " + str(question) + alternative)
            apply_style(cell, style_header)
            columnCounterHeader += 1
    rowCounter = 1
    columnCounterAnswers = columnCounter
#   comment Riet 12/5/2016: onderstaande nog te debuggen    
#    for participant in range(len(totalScore_loc)): # loop over participants
#        columnCounter = columnCounterAnswers;
#        serie = int(columnSeries_loc[participant]-1)
#        answers =  matrixAnswers[participant]
#        
#        #find questions with weight zero
#        #questionsZeroWeight = numpy.where(weightsQuestions_loc==0)
#                #find questions
#        #score[numpy.where(weightsQuestions_loc==0)[0]]=float('NaN')
#        
#        for question in range(1,numQuestions_loc+1):
#            questionInSerie = numpy.where(permutations_loc[serie]==question)[0][0]+1
#            answersQuestion = answers[numAlternatives_loc*(questionInSerie-1):numAlternatives_loc*(questionInSerie-1)+numAlternatives_loc]                                  
#            for counterAlternative in range(0,numAlternatives_loc):
#                print rowCounter,columnCounter,answersQuestion[counterAlternative]
#                sheetC.write(rowCounter,columnCounter,answersQuestion[counterAlternative])
#                columnCounter+=1
#        rowCounter+=1    


def write_participantsList(outputbook_loc, nameSheet_loc, deelnemers_loc):
    wb = outputbook_loc
    ws = wb.create_sheet(nameSheet_loc)

    columnCounter = 0
    rowCounter = 0
    ws.cell(row=rowCounter+1, column=columnCounter+1, value="1 tot 20")
    
    rowCounter += 1
    ws.cell(row=rowCounter+1, column=columnCounter+1, value="Naam")
    ws.cell(row=rowCounter+1, column=columnCounter+2, value="Voornaam")
    ws.cell(row=rowCounter+1, column=columnCounter+3, value="Studnr")
    ws.cell(row=rowCounter+1, column=columnCounter+11, value="TOTAAL")
    
    rowCounter += 1
    columnCounter = 2
    # Deelnemersnummers
    for i in range(len(deelnemers_loc)):
        ws.cell(row=rowCounter+1, column=columnCounter+1, value=deelnemers_loc[i])
        rowCounter += 1

def write_qsf(outputFolder_onderdeel_loc, numAlternatives_loc, numQuestions_loc, matrixAnswers_loc, correctAnswers_loc, deelnemers_loc, columnSeries_loc, jaar, toetsnaamOnderdeel, blankAnswer_loc):
    fqsf = open(outputFolder_onderdeel_loc + '/antwoorden_' + jaar + "_" + toetsnaamOnderdeel + '.qsf', 'w')
    fqsf.write('Snapshot,Participant,Vendor,Group')
    for question in range(1, numQuestions_loc + 1):
        fqsf.write(',Q' + str(question))
    fqsf.write('\n')
    
    # Replace letters by numbers for matrix answers and correct answers (sleutel)
    matrixAnswers_numbers = matrixAnswers_loc
    correctAnswers_numbers = correctAnswers_loc
    for alternative in range(numAlternatives_loc):
        letter = chr(97 + alternative).capitalize()
        matrixAnswers_numbers = np.where(matrixAnswers_numbers == letter, str(alternative + 1), matrixAnswers_numbers)
        correctAnswers_numbers = np.where(correctAnswers_numbers == letter, str(alternative + 1), correctAnswers_numbers)
    
    # The blank answers => replace by numAlternatives + 1
    letter = blankAnswer_loc
    matrixAnswers_numbers = np.where(matrixAnswers_numbers == letter, str(numAlternatives_loc + 1), matrixAnswers_numbers)

    # Write line with answers of participant
    for participant in range(len(deelnemers_loc)):
        fqsf.write(str(int(columnSeries_loc[participant])) + ',\"' + str(int(deelnemers_loc[participant])) + '\",\"Gravic, Inc.\",\"auto\"')
        antwoorden = matrixAnswers_numbers[participant]
        for vraag in range(numQuestions_loc):
            fqsf.write(',' + str(antwoorden[vraag]))
        fqsf.write('\n')
        
    # Write correct answers
    fqsf.write('1,\"999999\",\"Gravic, Inc.\",\"auto\"')
    for vraag in range(numQuestions_loc):
        fqsf.write(',' + str(correctAnswers_numbers[vraag]))
    fqsf.write('\n')
    
    # Close the file
    fqsf.close()