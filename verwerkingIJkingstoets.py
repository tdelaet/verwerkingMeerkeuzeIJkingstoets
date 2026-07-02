# -*- coding: utf-8 -*-
"""
Created on 9/7/2021

@author: tdelaet

assumptions:
    -There is a folder with name jaar_toets, e.g. 2021_ia21
     The folder contains:
        - OMR folder with:
            at least 1 excel file with OMR output:
                jaar_toets_OMRoutput_instelling.xslx e.g. 2021_ia21_OMRoutput_Brussel.xslx
        - onderdelen folder with:
            - maxScores_jaar_toets.txt (e.g. maxScores 2021_ia21.txt)
                that file contains the maximum scores for the whole test and the subparts
                 maxScore_TOTAAL, maxScore_A, maxScore_B, maxScore_C
                 bvb: 20,10,20,10 
                 conditions: (have to be integers separated by commas) en (length = 1 + number of subparts	)
            - sleutel_jaar_toets.txt (e.g. sleutel_2021_ia21.txt)
                that file contains the correct answers for the total test
                A,B,B,D,A,C ... 
                conditions: (have to be letters separated by commas) en (length = number of questions in total test	)
            - for each subpart a file jaar_toets_subpart.txt
                that contains the questions that are in that subpart
                e.g. 19,20,21,22
            - IF different series/permutations are used: permutatie_jaar_toets.txt (e.g. permutatie_2021_ia21.txt)
    - Below fill in variables jaar, toets, sessie, editie, aantal_onderdelen, numSeries, numAlternatives, blankAnswer, verwerking
"""

import string
import numpy
import matplotlib.pyplot as plt
import matplotlib.font_manager



plt.rcParams['font.family'] = 'DeJavu Serif'
plt.rcParams['font.serif'] = ['Times New Roman']

#from xlwt import Workbook
#import xlwt
#from xlrd import open_workbook
from openpyxl import load_workbook, Workbook
import os
import sys
import pandas as pd
import numpy as np


import checkInputVariables_openpyxl
import supportFunctions
import writeResults_openpyxl as writeResults
import leesSleutelEnPermutaties
import voorbereidingOnderdelen
import afwerkingOnderdelen
import plotFunctions
import warnings

#####################################################################################
#####################################################################################
### Variables to fill in
jaar = "2026"
sessie = 31
editie= "juli "+ jaar



#################################################
# =============================================================================
# toets = "ia" 
# aantal_onderdelen = 4 #TODO read from file or as extra safety?
# numSeries= 4 # number of series TODO lezen van file or as extra safety?
# neutralized=[] #%TODO: read from file or something else?
# instellingen = ["Leuven","Gent","Brussel","Kortrijk"] #ir en ia
# 
# =============================================================================
# #################################################
toets = "ir" 
aantal_onderdelen = 0 #todo read from file or as extra safety?
numSeries= 4 # number of series todo lezen van file or as extra safety?
neutralized=[] #%todo: read from file or something else?
instellingen = ["Leuven","Gent","Brussel","Kortrijk"] #ir en ia
#instellingen = ["Leuven","Kortrijk","Brussel","Gent"] #ir en ia

# #################################################
# toets = "ww" 
# aantal_onderdelen = 1 #TODO read from file or as extra safety?
# numSeries= 4 # number of series TODO lezen van file or as extra safety?
# neutralized=[] #%TODO: read from file or something else?
# instellingen = ["Leuven","Gent","Brussel","Kortrijk","Antwerpen","Hasselt"]


# #################################################
# toets = "ib" 
# aantal_onderdelen = 4 #TODO read from file or as extra safety?
# numSeries= 1 # number of series TODO lezen van file or as extra safety?
# neutralized=[] #%TODO: read from file or something else?
# instellingen = ["Leuven","Gent"]


#################################################
#toets = "ir" 
#aantal_onderdelen = 1 #todo read from file or as extra safety?
#numseries= 4 # number of series todo lezen van file or as extra safety?
#neutralized=[] #%todo: read from file or something else?
#instellingen = ["Leuven","Gent","Brussel","Kortrijk","Antwerpen","Hasselt","Antwerpen2","Leuven-Sophie"]

# #################################################
# toets = "bi" 
# aantal_onderdelen = 2 #TODO read from file or as extra safety?
# numSeries= 2 # number of series TODO lezen van file or as extra safety?
# neutralized=[] #%TODO: read from file or something else?
# instellingen = ["Leuven","Gent","Brussel","Kortrijk","Antwerpen"]

# #################################################
# toets = "fa" 
# aantal_onderdelen = 4 #TODO read from file or as extra safety?
# numSeries= 4 # number of series TODO lezen van file or as extra safety?
# neutralized=[] #%TODO: read from file or something else?
# instellingen = ["Leuven","Gent","Brussel","Kortrijk","Antwerpen"]


# #################################################
# toets = "in" 
# aantal_onderdelen = 4 #TODO read from file or as extra safety?
# numSeries= 2 # number of series TODO lezen van file or as extra safety?
# neutralized=[] #%TODO: read from file or something else?
# instellingen = ["Leuven","Leuven_Brugge", "Leuven_Geel", "Leuven_Gent", "Leuven_Sint-Katelijne-Waver","Gent","Gent_Kortrijk","Brussel","Antwerpen","Hasselt"]

#instellingen = ["all"] #ww
#instellingen = ["AE","BB","GR","LH","LK"] #wb
#instellingen = ["LZ"] #la
#instellingen = ["GC","LE"] #ib
#instellingen = ["LT"] #et
#instellingen = ["AE","BB","GD","LH","LK"] #bi
#instellingen = ["AE","BB","GR","LH","LK"] #wf
#instellingen = ["AE","BB","GZ","LH","LO","UH"] #rw
#instellingen = ["AE","BJ","GZ","LK","LZ","UH"] #bw
#instellingen = ["AE","BJ","GH","LK","LZ"] #fa
#instellingen = ["GT","LB","LK","LL","LN"] #hw
#instellingen = ["AA","BB","GT","LB","LK","LL","LN"] #hi
#instellingen = ["AA","AA_2","BB","GT","LB","LK","LL","LN"] #ew
#instellingen = ["AG","BB","CD","GB","GK","LE","LG","LO","LT","LW"] #in

# For actual rules see "afwerkingOnderdelen.py" bepaalGeslaagd en bepaalFeedbackGroep
if toets=="ia" or toets=="hw" or toets=="hi" or toets=="ew":
    regelFeedbackgroep = "ia"      #A als (TOTAAL >=maxTOTAAL/2 & scoreB>=maxScoreB/2)    
    regelGeslaagd = "ia"      #geslaagd als (TOTAAL >=maxTOTAAL/2 & scoreB>=maxScoreB/2)  
elif toets=="bi":
    regelFeedbackgroep="biib"  
    #feedbackgroep A score_Totaal>=12;
    #feedbackgroep B 10<=score_Totaal<12;
    #feedbackgroep C 5<=score_Totaal<10;
    #feedbackgroep D score_Totaal<5
    regelGeslaagd =  "geslaagdTotaal" #A als (TOTAAL >=maxTOTAAL/2) 
elif toets=="ib":
    regelFeedbackgroep="biib"  
    #feedbackgroep A score_A>=12;
    #feedbackgroep B 10<=score_A<12;
    #feedbackgroep C 5<=score_A<10;
    #feedbackgroep D score_A<5
    regelGeslaagd =  "ib" #A als (score_A >=10) 
elif toets=="bw":
    #feedbackgroep A score_TOTAAL >=10 
    #feedbackgroep B score_TOTAAL <10 AND score TOTAAL > 6
    #feedbackgroep C score_TOTAAL <=6
    regelGeslaagd =  "geslaagdTotaal" #A als (TOTAAL >=maxTOTAAL/2)  
    regelFeedbackgroep =  "bw"
elif  toets=="fa":
    #feedbackgroep A score_TOTAAL >=10 
    #feedbackgroep B score_TOTAAL <10 AND score TOTAAL > 7
    #feedbackgroep C score_TOTAAL <=7
    regelGeslaagd =  "geslaagdTotaal" #A als (TOTAAL >=maxTOTAAL/2)  
    regelFeedbackgroep =  "fa"
elif toets=="ir" or toets=="ww" or toets =="la":
    regelFeedbackgroep =  "irww" 
    #feedbackgroep A score_TOTAAL >=10 
    #feedbackgroep B score_TOTAAL <10 AND score TOTAAL > 5
    #feedbackgroep C score_TOTAAL <=5
    regelGeslaagd =  "geslaagdTotaal" #A als (TOTAAL >=maxTOTAAL/2) 
elif toets=="rw":
    regelFeedbackgroep =  "geslaagdTotaal" #A als (TOTAAL >=maxTOTAAL/2) 
    regelGeslaagd =  "geslaagdTotaal" #A als (TOTAAL >=maxTOTAAL/2) 
elif toets=="wf" or toets=="wb"   or  toets=="in" or toets=="et":
    regelFeedbackgroep =  "iedereenA"
    regelGeslaagd =  "geslaagdTotaal" #A als (TOTAAL >=maxTOTAAL/2) 
else:
    print ("ERROR found in input variables"   )
    sys.exit()

numAlternatives = 4 #number of alternatives


blankAnswer = "BLANK"  #how a blank answer is encoded in the OMR output
scoreBlankAnswer = -1.0/(float(numAlternatives)-1.0) #score for a blank answer
scoreWrongAnswer = -1.0/(float(numAlternatives)-1.0) # score for a wrong answer
scoreNeutralizedAnswer = 1.0

verwerking = "text" #als sleutel en permutatie als txt gegeven
#verwerking = "tex" #als sleutel en permutatie als tex zijn gegeven

toets = toets + str(sessie)
# do you want to write a feedback excel, one sheet per student?
writeFeedbackStudents = False


#####################################################################################
#####################################################################################

outputFolder = "../ijkingstoets-data/" + jaar + "/sessie " + str(sessie) + "/" + jaar + "_" +  toets
if not os.path.exists(outputFolder):
        os.makedirs(outputFolder)

# code from here
if numSeries==1:
     permutationsUsed = False
else:
     permutationsUsed = True

onderdelen = voorbereidingOnderdelen.voorbereidingOnderdelen(jaar,toets,sessie,permutationsUsed,aantal_onderdelen,instellingen,outputFolder,neutralized)
maxScores = numpy.loadtxt(outputFolder + "/onderdelen/maxScores_" + jaar+ "_"+ toets+ ".txt",delimiter=',',dtype="int",ndmin=1)


for onderdeel in (["TOTAAL"] + onderdelen):
    print("-------------------------------------------------")
    print("VERWERKING ONDERDEEL " + onderdeel)
    print("-------------------------------------------------")
    toetsnaamOnderdeel = toets + "_" + onderdeel
    folder_onderdeel = "../ijkingstoets-data/" + jaar + "/sessie " + str(sessie) + "/" + jaar + "_" +  toetsnaamOnderdeel
    # number of questions is length of the sleutel/correct answers
    numQuestions = len(numpy.loadtxt(folder_onderdeel + "/sleutel_"  + jaar +"_" + toetsnaamOnderdeel +".txt",delimiter=',',dtype="str",ndmin=1))
    print("aantal vragen "+ str(numQuestions))
    maxTotalScore = numpy.loadtxt(folder_onderdeel + "/maxScore_" + jaar+ "_"+ toetsnaamOnderdeel + ".txt",delimiter=',',dtype="int",ndmin=1)[0]
    print ("maximum score "+ str(maxTotalScore))
    
    
        
    #nameFile = "../OMR/test" #name of excel file with scanned forms
    nameSheet = "outputScan" #sheet name of excel file with scanned forms
    
    nameFile = folder_onderdeel + "/OMR/"+ jaar + "_" +  toetsnaamOnderdeel + "_OMRoutput" #name of excel file with scanned forms
    
    texinputFolder = folder_onderdeel + "/texinput/"        
    texoutputFolder = folder_onderdeel + "/texoutput/"
    if verwerking == "tex":
        if not os.path.exists(texoutputFolder):
            os.makedirs(texoutputFolder)    
    
    #where output of processing is saved
    outputFolder_onderdeel = folder_onderdeel + "/output_" + jaar + "_" + toetsnaamOnderdeel + "/"
    if not os.path.exists(outputFolder_onderdeel):
        os.makedirs(outputFolder_onderdeel)
    #where output for print en scan is save
    outputFolder_onderdeel_ps = folder_onderdeel + "/printenscan/"
    if not os.path.exists(outputFolder_onderdeel_ps):
        os.makedirs(outputFolder_onderdeel_ps)
    
    distributionList = [0.35,0.5,0.6,0.7,0.8,0.9]
    bordersDistributionStudentsLow =  [int(maxTotalScore*x) for x in distributionList]#for counting how many students get <=7,10 ...
    bordersDistributionStudentsHigh = bordersDistributionStudentsLow#for counting how many students get >=7,10 ...
    
    ############################
    #create list of expected content of scan file
    content = ["ijkID","vragenreeks"]
    
    for question in range(1,numQuestions+1):
            name = "Vraag" + str(question)
            content.append(name)
    ###########################

    ############################
    if verwerking == 'tex':
        #correct answers
        correctAnswers = leesSleutelEnPermutaties.leesSleutel(jaar,toetsnaamOnderdeel,texinputFolder)
        #permutations
        if numSeries == 1:
            permutations = numpy.zeros((1,numQuestions))
            for question in range(0,numQuestions):
                permutations[0,question] = question + 1
        else:
            permutations = leesSleutelEnPermutaties.leesPermutaties(jaar,toetsnaamOnderdeel,numSeries,texinputFolder)
    else:
        #correct answers
        correctAnswers = numpy.loadtxt(folder_onderdeel + "/sleutel_" + jaar+ "_"+ toetsnaamOnderdeel+ ".txt",delimiter=',',dtype="str",ndmin=1)
        #permutations
        if numSeries == 1:
            permutations = numpy.zeros((1,numQuestions))
            for question in range(0,numQuestions):
                permutations[0,question] = question + 1
        else:
            permutations = numpy.loadtxt(folder_onderdeel +  "/permutatie_" + jaar+ "_"+ toetsnaamOnderdeel+ ".txt",delimiter=',',dtype=float)
    
    neutralized_onderdeel=[]
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        neutralized_onderdeel = numpy.loadtxt(folder_onderdeel +  "/neutralized_" + jaar+ "_"+ toetsnaamOnderdeel+ ".txt",delimiter=',',dtype=float)
 
    print("sleutel: ")
    print(correctAnswers)
    print("permutaties: ")
    print(permutations)
    print("geneutralizeerd: ")
    print(neutralized_onderdeel)


    ### Important for tex processing, of texinputFolders does not exists, produces "empty" list
    #name of questions
    nameQuestions = leesSleutelEnPermutaties.leesNamenVragen(jaar,toetsnaamOnderdeel,texinputFolder,numQuestions)
    #classification of questions
    classificationQuestionsMod = leesSleutelEnPermutaties.leesClassificatieVragen(jaar,toetsnaamOnderdeel,texinputFolder,numQuestions)
    #categorie of questions
    categorieQuestions = leesSleutelEnPermutaties.leesCategorieVragen(jaar,toetsnaamOnderdeel,texinputFolder,numQuestions)

    #numpy.savetxt(outputFolder_onderdeel + "permutatie_"+ jaar +"_" + toetsnaamOnderdeel + ".txt",permutations,delimiter=',',fmt="%i")
    ############################
    ############################
    
    plt.close("all")

    #letters of answer alternatives
    alternatives = list(string.ascii_uppercase)[0:numAlternatives]
       
    if not( checkInputVariables_openpyxl.checkInputVariables(nameFile,nameSheet,numQuestions,numAlternatives,numSeries,correctAnswers,permutations,nameQuestions,instellingen,classificationQuestionsMod,categorieQuestions)):
        print ("ERROR found in input variables"   )
        sys.exit()
    print("after checkInputVariables_openpyxl")
        
    deelnemers_all = []      
    scoreQuestionsAllPermutations_all = []
    correctAnswersAllPermutations_all = []
    wrongAnswersAllPermutations_all = []
    blankAnswersAllPermutations_all = []
    neutralizedAnswersAllPermutations_all = []
    numQuestionsAlternatives_all = []
    scoreQuestionsIndicatedSeries_all = []
    numberCorrect_all = []
    numberCorrect_wrong = []
    numberCorrect_blank = []
    totalScoreDifferentPermutations_all = []
    totalScore_all = []
    averageScore_all = []
    medianScore_all = []
    standardDeviation_all = []
    percentagePass_all = []
    columnSeries_all = []
    matrixAnswers_all = []
    numParticipants_all = []
    scoreCategories_all = []
    numberCorrectAnswers_all = []
    numberWrongAnswers_all = []
    numberBlankAnswers_all = []
    numberNeutralizedAnswers_all = []    
    
    for instelling in instellingen:  
        counter = 0
        print ("instelling: " + instelling)
        #read file and get sheet
        # book= open_workbook(nameFile+"_"+ instelling+".xlsx") #TODO check if can set back to xslx (newer version python)
        # sheet = book.sheet_by_name(nameSheet)
            
        # #number of rows and columns
        # num_rows = sheet.nrows;
        # num_cols = sheet.ncols;
     
        # #number of participants = number of rows-1
        # numParticipants = num_rows-1;
        
        # #Load the first row => name indicating 
        # firstRow =  sheet.row(0) 
        # firstRowValues = sheet.row_values(0)
        
        # content_colNrs = supportFunctions.giveContentColNrs(content, sheet);
        
        # #prepare a matrix to store the score of a student with all possible permutations
        # scoreQuestionsIndicatedSeries= numpy.zeros((numParticipants,numQuestions))
        
        # # prepare output excels
        # outputbook = Workbook(style_compression=2)
        # outputbookperm = Workbook(style_compression=2)
        # outputStudentbook = Workbook(style_compression=2)
        # outputResults = Workbook(style_compression=2)
        
        # PYXLW Laad het bestaande Excel-bestand
        book = load_workbook(filename=nameFile + "_" + instelling + ".xlsx")
        sheet = book[nameSheet]
        
        # Aantal rijen en kolommen
        num_rows = sheet.max_row
        num_cols = sheet.max_column
        
        # Aantal deelnemers = aantal rijen - 1 (eerste rij is header)
        numParticipants = num_rows - 1
        
        # Eerste rij (kopteksten)
        firstRow = list(sheet.iter_rows(min_row=1, max_row=1))[0]
        firstRowValues = [cell.value for cell in firstRow]
        
        # Kolomnummers bepalen (je functie blijft hetzelfde)
        content_colNrs = supportFunctions.giveContentColNrs(content, sheet)
        
        # Matrix voorbereiden voor scores
        scoreQuestionsIndicatedSeries = np.zeros((numParticipants, numQuestions))
        
        # Nieuwe werkboeken aanmaken
        outputbook = Workbook()
        outputbookperm = Workbook()
        outputStudentbook = Workbook()
        outputResults = Workbook()

        
        
        if writeFeedbackStudents:
            outputFeedbackbook = Workbook(style_compression=2)
        
        name = "ijkID"
        studentenNrCol= content_colNrs[content.index(name)]
        #XLRD
        # deelnemers=sheet.col_values(studentenNrCol,1,num_rows)
        
        # if not supportFunctions.checkForUniqueParticipants(deelnemers):
        #     print ("ERROR: Duplicate participants found")
        #     sys.exit()
        
        # name = "vragenreeks"
        # #get the column in which the vragenreeks is stored
        # colNrSerie = content_colNrs[content.index(name)]
        # #get the series for the participants (so skip for row with name of first row)
        # columnSeries=sheet.col_values(colNrSerie,1,num_rows)
        
        #OPENPYXL
        # Get all rows from the sheet (excluding the header row)
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        num_rows = len(rows) + 1  # +1 because we skipped the header
        
        # Get the column of participant names (deelnemers)
        deelnemers = [row[studentenNrCol] for row in rows]
        
        if not supportFunctions.checkForUniqueParticipants(deelnemers):
            print("ERROR: Duplicate participants found")
            sys.exit()
        
        name = "vragenreeks"
        # Get the column index for "vragenreeks"
        colNrSerie = content_colNrs[content.index(name)]
        
        # Get the series for the participants
        columnSeries = [row[colNrSerie] for row in rows]

                  
        # get matrix of answers
        matrixAnswers = supportFunctions.getMatrixAnswers(sheet,content,correctAnswers,permutations,alternatives,numParticipants,columnSeries,content_colNrs)  
        supportFunctions.checkMatrixAnswers(matrixAnswers,alternatives,blankAnswer)
        
        #get the score for all permutations for each of the questions
        scoreQuestionsAllPermutations,correctAnswersAllPermutations,wrongAnswersAllPermutations,blankAnswersAllPermutations,neutralizedAnswersAllPermutations= supportFunctions.calculateScoreAllPermutations(sheet,blankAnswer,matrixAnswers,correctAnswers,permutations,alternatives,numParticipants,columnSeries,content_colNrs,scoreWrongAnswer,scoreBlankAnswer,scoreNeutralizedAnswer,neutralized_onderdeel)     
        #scoreQuestionsAllPermutations= supportFunctions.calculateScoreAllPermutations_old(sheet,content,correctAnswers,permutations,alternatives,numParticipants,columnSeries,content_colNrs)     
        numQuestionsAlternatives = supportFunctions.getNumberAlternatives(sheet,content,permutations,columnSeries,scoreQuestionsIndicatedSeries,alternatives,blankAnswer,content_colNrs)
        
        #get the scores for the indicated series
        scoreQuestionsIndicatedSeries, averageScoreQuestions, numberCorrectAnswers, numberWrongAnswers, numberBlankAnswers, numberNeutralizedAnswers =  supportFunctions.getScoreQuestionsIndicatedSeries(scoreQuestionsAllPermutations,correctAnswersAllPermutations,wrongAnswersAllPermutations,blankAnswersAllPermutations,neutralizedAnswersAllPermutations,columnSeries)
        
        #get the overall statistics
        totalScore, averageScore, medianScore, standardDeviation, percentagePass = supportFunctions.getOverallStatistics(scoreQuestionsIndicatedSeries,maxTotalScore)
        
        #get all the scores for the different permutations
        totalScoreDifferentPermutations = supportFunctions.calculateTotalScoreDifferentPermutations(scoreQuestionsAllPermutations,maxTotalScore)
        #print totalScoreDifferentPermutations
        
        #get the average score for the different question categories
        scoreCategories = supportFunctions.getScoreCategories(scoreQuestionsIndicatedSeries,categorieQuestions)
        
        numParticipantsSeries, averageScoreSeries, medianScoreSeries, standardDeviationSeries, percentagePassSeries, averageScoreQuestionsDifferentSeries = supportFunctions.getOverallStatisticsDifferentSeries(totalScoreDifferentPermutations,scoreQuestionsIndicatedSeries,columnSeries,maxTotalScore)
        
        totalScoreUpper,totalScoreMiddle,totalScoreLower,averageScoreUpper, averageScoreMiddle, averageScoreLower, averageScoreQuestionsUpper, averageScoreQuestionsMiddle, averageScoreQuestionsLower,numQuestionsAlternativesUpper,numQuestionsAlternativesMiddle,numQuestionsAlternativesLower, scoreQuestionsUpper, scoreQuestionsMiddle, scoreQuestionsLower,numUpper, numMiddle, numLower= supportFunctions.calculateUpperLowerStatistics(matrixAnswers,content,columnSeries,totalScore,scoreQuestionsIndicatedSeries,correctAnswers,alternatives,blankAnswer,content_colNrs,permutations)
         
        distributionStudentsHigh,distributionStudentsLow = supportFunctions.getDistributionStudents(totalScore,bordersDistributionStudentsLow,bordersDistributionStudentsHigh)
        
        #only plot and write to files if there are multiple instellingen. Otherwise "geheel" is same as single instelling, producing double output.
        if (len(instellingen)!=1):
            ## WRITING THE OUTPUT TO A FILE
            writeResults.write_results(outputbook,outputbookperm,numQuestions,correctAnswers,alternatives,blankAnswer,
                              maxTotalScore,content,content_colNrs,
                              columnSeries,deelnemers,
                              numParticipants,
                              totalScore,percentagePass,
                              scoreQuestionsIndicatedSeries,
                              totalScoreDifferentPermutations,
                              medianScore,
                              standardDeviation,
                              averageScore,averageScoreUpper,averageScoreMiddle,averageScoreLower,
                              averageScoreQuestions,averageScoreQuestionsUpper,averageScoreQuestionsMiddle,averageScoreQuestionsLower,
                              averageScoreQuestionsDifferentSeries,
                              numUpper,numMiddle,numLower,
                              numParticipantsSeries,
                              averageScoreSeries,medianScoreSeries,standardDeviationSeries,percentagePassSeries,
                              numQuestionsAlternatives, numQuestionsAlternativesUpper, numQuestionsAlternativesMiddle, numQuestionsAlternativesLower,
                              nameQuestions,classificationQuestionsMod,categorieQuestions,
                              bordersDistributionStudentsLow,bordersDistributionStudentsHigh,
                              distributionStudentsLow,distributionStudentsHigh)
            
                         
            ## WRITING A FILE TO UPLOAD TO TOLEDO WITH THE GRADES
            writeResults.write_scoreStudents(outputStudentbook,"punten",permutations,numParticipants,deelnemers, numQuestions,numAlternatives,content,content_colNrs,totalScore,scoreQuestionsIndicatedSeries,columnSeries,matrixAnswers,numberCorrectAnswers, numberWrongAnswers, numberBlankAnswers,numberNeutralizedAnswers)           
            #writeResults.write_resultsFile(outputResults,"resultaten",permutations,numParticipants,deelnemers, numQuestions,numAlternatives,content,content_colNrs,totalScore,scoreQuestionsIndicatedSeries,columnSeries,matrixAnswers,numberCorrectAnswers, numberWrongAnswers, numberBlankAnswers)                   

            writeResults.write_scoreCategoriesStudents(outputStudentbook,"percentageCategorien",deelnemers, totalScore, categorieQuestions, scoreCategories)
            
            # outputFolder_instelling = outputFolder_onderdeel + instelling + "/"
            # if not os.path.exists(outputFolder_instelling):
            #     os.makedirs(outputFolder_instelling)    
            # outputbook.save(outputFolder_instelling + 'output_'  + jaar + "_" +  toets + "_" +instelling+'.xlsx') 
            # outputbookperm.save(outputFolder_instelling + 'output_permutations_'  + jaar + "_" +  toets + "_" +instelling+'.xlsx') 
            # outputStudentbook.save(outputFolder_instelling + 'punten_'  + jaar + "_" +  toets + "_" +instelling+'.xlsx') 
            
            # Create the output folder if it doesn't exist
            outputFolder_instelling = outputFolder_onderdeel + instelling + "/"
            if not os.path.exists(outputFolder_instelling):
                os.makedirs(outputFolder_instelling)
            
            # Save the workbooks
            outputbook.save(outputFolder_instelling + 'output_' + jaar + "_" + toets + "_" + instelling + '.xlsx')
            outputbookperm.save(outputFolder_instelling + 'output_permutations_' + jaar + "_" + toets + "_" + instelling + '.xlsx')
            outputStudentbook.save(outputFolder_instelling + 'punten_' + jaar + "_" + toets + "_" + instelling + '.xlsx')


        deelnemers_all.append(deelnemers)
        scoreQuestionsAllPermutations_all.append(scoreQuestionsAllPermutations)
        correctAnswersAllPermutations_all.append(correctAnswersAllPermutations)
        wrongAnswersAllPermutations_all.append(wrongAnswersAllPermutations)
        blankAnswersAllPermutations_all.append(blankAnswersAllPermutations)
        neutralizedAnswersAllPermutations_all.append(neutralizedAnswersAllPermutations)
        numQuestionsAlternatives_all.append(numQuestionsAlternatives)
        scoreQuestionsIndicatedSeries_all.append(scoreQuestionsIndicatedSeries)
        totalScoreDifferentPermutations_all.append(totalScoreDifferentPermutations)
        totalScore_all.append(totalScore)
        averageScore_all.append(averageScore)
        medianScore_all.append(medianScore)
        standardDeviation_all.append(standardDeviation)
        percentagePass_all.append(percentagePass)
        columnSeries_all.append(columnSeries)
        matrixAnswers_all.append(matrixAnswers)
        numParticipants_all.append(numParticipants)
        scoreCategories_all.append(scoreCategories)
        numberCorrectAnswers_all.append(numberCorrectAnswers)
        numberWrongAnswers_all.append(numberWrongAnswers)
        numberBlankAnswers_all.append(numberBlankAnswers)
        numberNeutralizedAnswers_all.append(numberNeutralizedAnswers)

    
    deelnemers_tot = numpy.hstack(deelnemers_all)
    scoreQuestionsAllPermutations_tot = numpy.hstack(scoreQuestionsAllPermutations_all)
    numQuestionsAlternatives_tot = sum(numQuestionsAlternatives_all)
    scoreQuestionsIndicatedSeries_tot = numpy.vstack(scoreQuestionsIndicatedSeries_all)
    totalScoreDifferentPermutations_tot = numpy.vstack(totalScoreDifferentPermutations_all)
    correctAnswersAllPermutations_tot = numpy.hstack(correctAnswersAllPermutations_all)
    wrongAnswersAllPermutations_tot = numpy.hstack(wrongAnswersAllPermutations_all)
    blankAnswersAllPermutations_tot = numpy.hstack(blankAnswersAllPermutations_all)
    neutralizedAnswersAllPermutations_tot = numpy.hstack(neutralizedAnswersAllPermutations_all)        
    columnSeries_tot = numpy.hstack(columnSeries_all)
    matrixAnswers_tot = numpy.vstack(matrixAnswers_all)
    numParticipants_stacked_tot = numpy.vstack(numParticipants_all)
    averageScore_stacked_tot = numpy.vstack(averageScore_all)
    medianScore_stacked_tot = numpy.vstack(medianScore_all)
    standardDeviation_stacked_tot = numpy.vstack(standardDeviation_all)
    percentagePass_stacked_tot  = numpy.vstack(percentagePass_all)
    numParticipants_tot = sum(numParticipants_stacked_tot)[0]
    scoreCategories_tot = numpy.hstack(scoreCategories_all)
    numberCorrectAnswers_tot = numpy.hstack(numberCorrectAnswers_all)
    numberWrongAnswers_tot = numpy.hstack(numberWrongAnswers_all)
    numberBlankAnswers_tot = numpy.hstack(numberBlankAnswers_all)
    numberNeutralizedAnswers_tot = numpy.hstack(numberNeutralizedAnswers_all)
    
    totalScore_tot, averageScore_tot, medianScore_tot, standardDeviation_tot, percentagePass_tot = supportFunctions.getOverallStatistics(scoreQuestionsIndicatedSeries_tot,maxTotalScore)
    numParticipantsSeries_tot, averageScoreSeries_tot, medianScoreSeries_tot, standardDeviationSeries_tot, percentagePassSeries_tot, averageScoreQuestionsDifferentSeries_tot = supportFunctions.getOverallStatisticsDifferentSeries(totalScoreDifferentPermutations_tot,scoreQuestionsIndicatedSeries_tot,columnSeries_tot,maxTotalScore)
    scoreQuestionsIndicatedSeries_tot, averageScoreQuestions_tot, numberCorrectAnswers_tot, numberWrongAnswers_tot, numberBlankAnswers_tot, numberNeutralizedAnswers_tot =  supportFunctions.getScoreQuestionsIndicatedSeries(scoreQuestionsAllPermutations_tot,correctAnswersAllPermutations_tot,wrongAnswersAllPermutations_tot,blankAnswersAllPermutations_tot,neutralizedAnswersAllPermutations_tot,columnSeries_tot)
        
    totalScoreUpper_tot,totalScoreMiddle_tot,totalScoreLower_tot,averageScoreUpper_tot, averageScoreMiddle_tot, averageScoreLower_tot, averageScoreQuestionsUpper_tot, averageScoreQuestionsMiddle_tot, averageScoreQuestionsLower_tot,numQuestionsAlternativesUpper_tot,numQuestionsAlternativesMiddle_tot,numQuestionsAlternativesLower_tot, scoreQuestionsUpper_tot, scoreQuestionsMiddle_tot, scoreQuestionsLower_tot,numUpper_tot, numMiddle_tot, numLower_tot= supportFunctions.calculateUpperLowerStatistics(matrixAnswers_tot,content,columnSeries_tot,totalScore_tot,scoreQuestionsIndicatedSeries_tot,correctAnswers,alternatives,blankAnswer,content_colNrs,permutations)
    distributionStudentsHigh_tot,distributionStudentsLow_tot= supportFunctions.getDistributionStudents(totalScore_tot,bordersDistributionStudentsLow,bordersDistributionStudentsHigh)

    
    # # write to excel_file
    # outputbook = xlwt.Workbook(style_compression=2)
    # outputbookperm = xlwt.Workbook(style_compression=2)
    # outputStudentbook = xlwt.Workbook(style_compression=2)  
    # outputResults = xlwt.Workbook(style_compression=2)  
    # outputInstellingen = xlwt.Workbook(style_compression=2)  
    # if writeFeedbackStudents:
    #     outputFeedbackbook = xlwt.Workbook(style_compression=2)
    # outputFeedbackPlatformbook = xlwt.Workbook(style_compression=2)
    # outputDeelnemersLijst = xlwt.Workbook(style_compression=2)
    # Nieuwe werkboeken aanmaken
    outputbook = Workbook()
    outputbookperm = Workbook()
    outputStudentbook = Workbook()
    outputResults = Workbook()
    outputInstellingen = Workbook() 
    if writeFeedbackStudents:
         outputFeedbackbook = Workbook()
    outputFeedbackPlatformbook = Workbook()
    outputDeelnemersLijst = Workbook()
    

    ## WRITING THE OUTPUT TO A FILE
    writeResults.write_qsf(outputFolder_onderdeel_ps,numAlternatives,numQuestions,matrixAnswers_tot,correctAnswers,deelnemers_tot,columnSeries_tot,jaar,toetsnaamOnderdeel,blankAnswer)
    writeResults.write_results(outputbook,outputbookperm,numQuestions,correctAnswers,alternatives,blankAnswer,
                      maxTotalScore,content,content_colNrs,
                      columnSeries_tot,deelnemers_tot,
                      numParticipants_tot,
                      totalScore_tot,percentagePass_tot,
                      scoreQuestionsIndicatedSeries_tot,
                      totalScoreDifferentPermutations_tot,
                      medianScore_tot,
                      standardDeviation_tot,
                      averageScore_tot,averageScoreUpper_tot,averageScoreMiddle_tot,averageScoreLower_tot,
                      averageScoreQuestions_tot,averageScoreQuestionsUpper_tot,averageScoreQuestionsMiddle_tot,averageScoreQuestionsLower_tot,
                      averageScoreQuestionsDifferentSeries_tot,
                      numUpper_tot,numMiddle_tot,numLower_tot,
                      numParticipantsSeries_tot,
                      averageScoreSeries_tot,medianScoreSeries_tot,standardDeviationSeries_tot,percentagePassSeries_tot,
                      numQuestionsAlternatives_tot, numQuestionsAlternativesUpper_tot, numQuestionsAlternativesMiddle_tot, numQuestionsAlternativesLower_tot,
                      nameQuestions,classificationQuestionsMod,categorieQuestions,
                      bordersDistributionStudentsLow,bordersDistributionStudentsHigh,distributionStudentsLow_tot,distributionStudentsHigh_tot
                      )    
    
    writeResults.write_scoreStudents(outputStudentbook,"punten",permutations,numParticipants_tot,deelnemers_tot, numQuestions,numAlternatives,content,content_colNrs,totalScore_tot,scoreQuestionsIndicatedSeries_tot,columnSeries_tot,matrixAnswers_tot,numberCorrectAnswers_tot,numberWrongAnswers_tot,numberBlankAnswers_tot,numberNeutralizedAnswers_tot)            
    writeResults.write_resultsFile(outputResults,"resultaten",permutations,numParticipants_tot,deelnemers_tot, numQuestions,numAlternatives,content,content_colNrs,totalScore_tot,scoreQuestionsIndicatedSeries_tot,columnSeries_tot,matrixAnswers_tot,numberCorrectAnswers_tot,numberWrongAnswers_tot,numberBlankAnswers_tot,numberNeutralizedAnswers_tot)           
    writeResults.write_overallStatisticsInstellingen(outputInstellingen,"instellingen",instellingen,numParticipants_tot,numParticipants_stacked_tot,averageScore_tot,averageScore_stacked_tot,medianScore_tot,medianScore_stacked_tot,standardDeviation_tot,standardDeviation_stacked_tot,percentagePass_tot,percentagePass_stacked_tot)
    writeResults.write_scoreStudentsNonPermutated(outputStudentbook,"punten_reeks1",permutations,numParticipants,deelnemers, numQuestions,numAlternatives,alternatives,content,content_colNrs,totalScore,scoreQuestionsIndicatedSeries,columnSeries,matrixAnswers)
    writeResults.write_scoreCategoriesStudents(outputStudentbook,"percentageCategorien",deelnemers_tot,totalScore_tot, categorieQuestions, scoreCategories_tot)
    

    if writeFeedbackStudents:
        writeResults.write_feedbackStudents(outputFeedbackbook,permutations,numParticipants_tot,deelnemers_tot, numQuestions,
                                        alternatives,numAlternatives,content,content_colNrs,
                                        totalScore_tot,scoreQuestionsIndicatedSeries_tot,columnSeries_tot,matrixAnswers_tot,
                                        categorieQuestions,scoreCategories_tot,
                                        averageScoreQuestions_tot,averageScoreQuestionsUpper_tot,averageScoreQuestionsMiddle_tot,averageScoreQuestionsLower_tot
                                        ,correctAnswers, numQuestionsAlternatives_tot)
    writeResults.write_participantsList(outputDeelnemersLijst,"Beoordelingen",deelnemers_tot)

    outputbook.save(outputFolder_onderdeel + 'output' +'_controleerVoorKwaliteitToets_'  + jaar + "_" +  toetsnaamOnderdeel + '.xlsx')  
    outputbookperm.save(outputFolder_onderdeel + 'output_controleerVoorFouteReeksen_'  + jaar + "_" +  toetsnaamOnderdeel + '.xlsx')  
    if (len(instellingen)!=1):
        outputInstellingen.save(outputFolder_onderdeel + 'instellingen_'  + jaar + "_" +  toetsnaamOnderdeel + '.xlsx')  
    outputStudentbook.save(outputFolder_onderdeel + 'punten_'  + jaar + "_" +  toetsnaamOnderdeel + '.xlsx')  
    outputResults.save(outputFolder_onderdeel + '../resultaten_'  + jaar + "_" +  toetsnaamOnderdeel + '.xlsx')  
    outputDeelnemersLijst.save(outputFolder_onderdeel_ps + 'deelnemerslijst_KULoket_'  + jaar + "_" +  toetsnaamOnderdeel + '.xlsx')  
    if writeFeedbackStudents:
        outputFeedbackbook.save(outputFolder_onderdeel+ 'feedback_'  + jaar + "_" +  toetsnaamOnderdeel + '.xlsx')  

   
punten_compose,geslaagdVariabele = afwerkingOnderdelen.genereerPuntenBestand(jaar,toets,sessie,onderdelen,regelFeedbackgroep,regelGeslaagd,maxScores,outputFolder)
afwerkingOnderdelen.kopieerQSF(jaar,toets,outputFolder)

teller=0
for onderdeel in (["TOTAAL"] + onderdelen):
    print("-------------------------------------------------")
    print("Plotting ONDERDEEL " + onderdeel)
    print("-------------------------------------------------")
    toetsnaamOnderdeel = toets + "_" + onderdeel
    if onderdeel == "TOTAAL":
        totalScore_tot=punten_compose[onderdeel]
        percentagePassed = 100*sum(geslaagdVariabele=="True")/float(len(totalScore_tot))
    else:
        totalScore_tot=punten_compose["score"+onderdeel]
        percentagePassed = 100*sum(score>= maxScores[teller]/2.0 for score in totalScore_tot)/float(len(totalScore_tot) )
    print("percentage passed " + str(int(percentagePassed)))
    folder_onderdeel = "../ijkingstoets-data/" + jaar + "/sessie " + str(sessie) + "/" + jaar + "_" +  toetsnaamOnderdeel  
    #where output of processing is saved
    outputFolder_onderdeel = folder_onderdeel + "/output_" + jaar + "_" + toetsnaamOnderdeel + "/"
    # plot the histogram of the total score
    saveNameFig = outputFolder_onderdeel + 'histogramGeheel_'  + jaar + "_" +  toetsnaamOnderdeel + '.png'
    plotFunctions.plotHistogram(saveNameFig,maxScores[teller],totalScore_tot,percentagePassed)
    
    saveNameFig = outputFolder_onderdeel + 'verdelingDeelnemers_'  + jaar + "_" +  toetsnaamOnderdeel + '.png'
    plotFunctions.plotPieParticipants(saveNameFig,instellingen,numParticipants_all) 
    teller=teller+1
    
# Let op, zips creëren duurt wel even
afwerkingOnderdelen.genereerZIPs(jaar,toets,sessie,onderdelen,outputFolder)